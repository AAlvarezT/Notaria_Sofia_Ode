#!/usr/bin/env python3
"""Conciliador de comprobantes pendientes: SIGNO vs. Créditos actuales.

Uso rápido:
  1) Coloque este archivo en la misma carpeta que:
       - CREDITOS SISTEMA SIGNO.xlsx
       - CREDITOS ACTUALES 2026.xlsx
  2) Ejecute: python actualizar_creditos.py

También admite rutas explícitas:
  python actualizar_creditos.py --signo "reporte.xlsx" --creditos "creditos.xlsx" --salida "resultado.xlsx"

Requisito:
  pip install openpyxl

Versión de máxima compatibilidad con Excel de escritorio.
No usa tablas estructuradas ni formato condicional, evitando reparaciones de styles.xml.
"""

from __future__ import annotations

import argparse
import re
import sys
import unicodedata
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: falta openpyxl. Instálelo con: pip install openpyxl")
    input("Presione Enter para cerrar...")
    raise SystemExit(1)

DEFAULT_SIGNO = "CREDITOS SISTEMA SIGNO.xlsx"
DEFAULT_CREDITOS = "CREDITOS ACTUALES 2026.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
RED_FILL = PatternFill("solid", fgColor="F4CCCC")
YELLOW_FILL = PatternFill("solid", fgColor="FFF2CC")
GREEN_FILL = PatternFill("solid", fgColor="D9EAD3")
GRAY_FILL = PatternFill("solid", fgColor="E7E6E6")
WHITE_FONT = Font(color="FFFFFF", bold=True)
THIN_GRAY = Side(style="thin", color="D9E1F2")


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    text = re.sub(r"\s+", " ", text)
    return text


def normalize_text(value: Any) -> str:
    text = clean_text(value).upper()
    text = unicodedata.normalize("NFKD", text)
    return "".join(c for c in text if not unicodedata.combining(c))


def normalize_ruc(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\D", "", str(value))


def normalize_doc(value: Any, width: int = 7) -> str:
    """Normaliza 30025 y 0030025 al mismo documento de 7 dígitos."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    raw = clean_text(value)
    digits = re.sub(r"\D", "", raw)
    return digits.zfill(width) if digits else normalize_text(raw)


def parse_date(value: Any) -> Optional[datetime]:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    text = clean_text(value)
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def to_number(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = clean_text(value).replace("S/", "").replace(",", "")
    try:
        return float(text)
    except ValueError:
        return 0.0


def make_key(tipo: Any, serie: Any, documento: Any) -> str:
    t = normalize_text(tipo)
    s = normalize_text(serie)
    d = normalize_doc(documento)
    return f"{t}|{s}|{d}" if t and s and d else ""


def locate_header(ws, required: Iterable[str], max_rows: int = 20, max_cols: int = 30) -> Tuple[int, Dict[str, int]]:
    wanted = {normalize_text(x): x for x in required}
    for row in range(1, min(ws.max_row, max_rows) + 1):
        found: Dict[str, int] = {}
        for col in range(1, min(ws.max_column, max_cols) + 1):
            normalized = normalize_text(ws.cell(row, col).value)
            if normalized in wanted:
                found[wanted[normalized]] = col
        if len(found) >= max(3, len(required) // 2):
            return row, found
    raise ValueError(f"No se encontró la fila de encabezados en la hoja {ws.title!r}.")


def read_signo(path: Path) -> List[Dict[str, Any]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    required = ["RUC", "GRUPO", "EMISION", "TIPO", "SERIE A", "DOC", "CLIENTE", "KARDEX", "FACTURADO", "SALDO", "VENCIMIENTO", "EMPLEADO", "COMENTARIOS"]
    header_row, cols = locate_header(ws, required)
    records: List[Dict[str, Any]] = []
    for r in range(header_row + 1, ws.max_row + 1):
        tipo = ws.cell(r, cols["TIPO"]).value
        serie = ws.cell(r, cols["SERIE A"]).value
        doc = ws.cell(r, cols["DOC"]).value
        key = make_key(tipo, serie, doc)
        if not key:
            continue
        records.append({
            "ID": key,
            "RUC": normalize_ruc(ws.cell(r, cols["RUC"]).value),
            "GRUPO": clean_text(ws.cell(r, cols["GRUPO"]).value),
            "CLIENTE": clean_text(ws.cell(r, cols["CLIENTE"]).value),
            "EMISION": parse_date(ws.cell(r, cols["EMISION"]).value),
            "TIPO": normalize_text(tipo),
            "SERIE": normalize_text(serie),
            "DOC": normalize_doc(doc),
            "TRAMITE": clean_text(ws.cell(r, cols["KARDEX"]).value),
            "FACTURADO": to_number(ws.cell(r, cols["FACTURADO"]).value),
            "SALDO": to_number(ws.cell(r, cols["SALDO"]).value),
            "VENCIMIENTO": parse_date(ws.cell(r, cols["VENCIMIENTO"]).value),
            "EMPLEADO": clean_text(ws.cell(r, cols["EMPLEADO"]).value),
            "COMENTARIO_SIGNO": clean_text(ws.cell(r, cols["COMENTARIOS"]).value),
        })
    wb.close()
    return records


def read_creditos(path: Path) -> List[Dict[str, Any]]:
    wb = load_workbook(path, data_only=True, read_only=False)
    ws = wb["CREDITOS"] if "CREDITOS" in wb.sheetnames else wb.active
    aliases = {
        "RUC/DNI": "RUC/DNI", "RAZON SOCIAL": "RAZON SOCIAL", "EMISION": "EMISION",
        "TIPO": "TIPO", "SERIE": "SERIE", "DOC": "DOC", "TRAMITE": "TRAMITE",
        "PENDIENTE DE PAGO NOTARIAL": "NOTARIAL", "PENDIENTE DE PAGO REGISTRAL": "REGISTRAL",
        "PENDIENTE DE PAGO GESTION": "GESTION", "VENCIMIENTO": "VENCIMIENTO",
        "DETALLE DE COBRANZA": "DETALLE", "SOLICITANTE": "SOLICITANTE"
    }
    header_row = None
    cols: Dict[str, int] = {}
    for r in range(1, min(ws.max_row, 20) + 1):
        temp = {}
        for c in range(1, min(ws.max_column, 30) + 1):
            n = normalize_text(ws.cell(r, c).value)
            if n in aliases:
                temp[aliases[n]] = c
        if all(k in temp for k in ("TIPO", "SERIE", "DOC")):
            header_row, cols = r, temp
            break
    if header_row is None:
        raise ValueError("No se encontraron encabezados TIPO, SERIE y DOC en Créditos actuales.")

    records: List[Dict[str, Any]] = []
    for r in range(header_row + 1, ws.max_row + 1):
        tipo = ws.cell(r, cols["TIPO"]).value
        serie = ws.cell(r, cols["SERIE"]).value
        doc = ws.cell(r, cols["DOC"]).value
        key = make_key(tipo, serie, doc)
        if not key:
            continue  # omite títulos, subtotales y filas vacías
        notarial = to_number(ws.cell(r, cols.get("NOTARIAL", 0)).value) if cols.get("NOTARIAL") else 0
        registral = to_number(ws.cell(r, cols.get("REGISTRAL", 0)).value) if cols.get("REGISTRAL") else 0
        gestion = to_number(ws.cell(r, cols.get("GESTION", 0)).value) if cols.get("GESTION") else 0
        records.append({
            "ID": key,
            "RUC": normalize_ruc(ws.cell(r, cols.get("RUC/DNI", 0)).value) if cols.get("RUC/DNI") else "",
            "CLIENTE": clean_text(ws.cell(r, cols.get("RAZON SOCIAL", 0)).value) if cols.get("RAZON SOCIAL") else "",
            "EMISION": parse_date(ws.cell(r, cols.get("EMISION", 0)).value) if cols.get("EMISION") else None,
            "TIPO": normalize_text(tipo), "SERIE": normalize_text(serie), "DOC": normalize_doc(doc),
            "TRAMITE": clean_text(ws.cell(r, cols.get("TRAMITE", 0)).value) if cols.get("TRAMITE") else "",
            "NOTARIAL": notarial, "REGISTRAL": registral, "GESTION": gestion,
            "SALDO_MANUAL": notarial + registral + gestion,
            "VENCIMIENTO": parse_date(ws.cell(r, cols.get("VENCIMIENTO", 0)).value) if cols.get("VENCIMIENTO") else None,
            "DETALLE_COBRANZA": clean_text(ws.cell(r, cols.get("DETALLE", 0)).value) if cols.get("DETALLE") else "",
            "SOLICITANTE": clean_text(ws.cell(r, cols.get("SOLICITANTE", 0)).value) if cols.get("SOLICITANTE") else "",
        })
    wb.close()
    return records


def deduplicate(records: List[Dict[str, Any]], source: str) -> Tuple[Dict[str, Dict[str, Any]], List[str]]:
    result: Dict[str, Dict[str, Any]] = {}
    duplicates: List[str] = []
    for record in records:
        key = record["ID"]
        if key in result:
            duplicates.append(f"{source}: {key}")
        result[key] = record
    return result, duplicates


def classify(vencimiento: Optional[datetime], saldo: float, today: date) -> Tuple[str, Optional[int]]:
    if saldo <= 0:
        return "SIN SALDO", None
    if vencimiento is None:
        return "SIN VENCIMIENTO", None
    days = (vencimiento.date() - today).days
    if days < -30: return "CRÍTICA", days
    if days < 0: return "VENCIDA", days
    if days == 0: return "VENCE HOY", days
    if days <= 3: return "ALTA", days
    if days <= 7: return "PREVENTIVA", days
    return "AL DÍA", days


def reconcile(signo_records: List[Dict[str, Any]], old_records: List[Dict[str, Any]]) -> Dict[str, Any]:
    signo, dup_signo = deduplicate(signo_records, "SIGNO")
    old, dup_old = deduplicate(old_records, "CRÉDITOS")
    signo_keys, old_keys = set(signo), set(old)
    new_keys = sorted(signo_keys - old_keys)
    missing_keys = sorted(old_keys - signo_keys)
    common_keys = sorted(signo_keys & old_keys)
    today = date.today()

    master = []
    changes = []
    for key in sorted(signo_keys | old_keys):
        s = signo.get(key)
        o = old.get(key)
        if s and o:
            status = "PENDIENTE EN SIGNO"
            saldo_previous, saldo_current = o["SALDO_MANUAL"], s["SALDO"]
            if abs(saldo_previous - saldo_current) > 0.009:
                changes.append({"ID": key, "CLIENTE": s["CLIENTE"] or o["CLIENTE"], "SALDO_ANTERIOR": saldo_previous, "SALDO_ACTUAL": saldo_current, "DIFERENCIA": saldo_current - saldo_previous})
        elif s:
            status, saldo_previous, saldo_current = "NUEVO EN SIGNO", 0.0, s["SALDO"]
        else:
            status, saldo_previous, saldo_current = "VALIDAR PAGO/ANULACIÓN", o["SALDO_MANUAL"], 0.0
        source = s or o
        priority, days = classify((s or o).get("VENCIMIENTO"), saldo_current if s else saldo_previous, today)
        master.append({
            "ID": key, "RUC": (s or {}).get("RUC") or (o or {}).get("RUC", ""),
            "CLIENTE": (s or {}).get("CLIENTE") or (s or {}).get("GRUPO") or (o or {}).get("CLIENTE", ""),
            "TIPO": source.get("TIPO", ""), "SERIE": source.get("SERIE", ""), "DOC": source.get("DOC", ""),
            "EMISION": source.get("EMISION"), "TRAMITE": source.get("TRAMITE", ""),
            "FACTURADO_SIGNO": (s or {}).get("FACTURADO", 0.0), "SALDO_SIGNO": (s or {}).get("SALDO", 0.0),
            "SALDO_MANUAL": (o or {}).get("SALDO_MANUAL", 0.0), "VENCIMIENTO": source.get("VENCIMIENTO"),
            "ESTADO_CONCILIACION": status, "PRIORIDAD": priority, "DIAS_PARA_VENCER": days,
            "DETALLE_COBRANZA": (o or {}).get("DETALLE_COBRANZA", ""),
            "SOLICITANTE": (o or {}).get("SOLICITANTE", "") or (s or {}).get("EMPLEADO", ""),
            "ULTIMA_SINCRONIZACION": datetime.now(),
        })

    priority_order = {"CRÍTICA": 0, "VENCIDA": 1, "VENCE HOY": 2, "ALTA": 3, "PREVENTIVA": 4, "SIN VENCIMIENTO": 5, "AL DÍA": 6, "SIN SALDO": 7}
    collections = sorted([x for x in master if x["ESTADO_CONCILIACION"] != "VALIDAR PAGO/ANULACIÓN" and x["SALDO_SIGNO"] > 0], key=lambda x: (priority_order.get(x["PRIORIDAD"], 99), x["DIAS_PARA_VENCER"] if x["DIAS_PARA_VENCER"] is not None else 99999, -x["SALDO_SIGNO"]))
    return {
        "master": master, "new": [signo[k] for k in new_keys], "possible_paid": [old[k] for k in missing_keys],
        "changes": changes, "collections": collections, "duplicates": dup_signo + dup_old,
        "metrics": {
            "Pendientes en SIGNO": len(signo), "Registros válidos en Créditos actuales": len(old),
            "Nuevos en SIGNO": len(new_keys), "Posibles pagos/anulaciones": len(missing_keys),
            "Saldos modificados": len(changes), "Saldo total pendiente SIGNO": sum(x["SALDO"] for x in signo.values()),
            "Comprobantes vencidos": sum(1 for x in signo.values() if x.get("VENCIMIENTO") and x["VENCIMIENTO"].date() < date.today() and x["SALDO"] > 0),
        }
    }


def add_table_sheet(wb: Workbook, title: str, rows: List[Dict[str, Any]], columns: List[Tuple[str, str]], table_name: str) -> None:
    """Crea una hoja tabular compatible, sin objetos Table de OOXML."""
    ws = wb.create_sheet(title)
    headers = [label for _, label in columns]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in rows:
        ws.append([row.get(key) for key, _ in columns])
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows)+1}"
    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.border = Border(bottom=Side(style="thin", color="D9E1F2"))
    for idx, (key, _) in enumerate(columns, start=1):
        values = [headers[idx - 1]] + [clean_text(r.get(key)) for r in rows[:200]]
        width = min(max(len(v) for v in values) + 2, 38)
        ws.column_dimensions[get_column_letter(idx)].width = max(width, 11)
        column_cells = list(ws[get_column_letter(idx)])[1:]
        if key in {"EMISION", "VENCIMIENTO"}:
            for cell in column_cells:
                cell.number_format = "dd/mm/yyyy"
        elif key == "ULTIMA_SINCRONIZACION":
            for cell in column_cells:
                cell.number_format = "dd/mm/yyyy hh:mm"
        elif key in {"FACTURADO", "SALDO", "FACTURADO_SIGNO", "SALDO_SIGNO", "SALDO_MANUAL", "SALDO_ANTERIOR", "SALDO_ACTUAL", "DIFERENCIA", "NOTARIAL", "REGISTRAL", "GESTION"}:
            # Formato estándar para evitar corrupción de styles.xml.
            for cell in column_cells:
                cell.number_format = "#,##0.00"


def write_output(result: Dict[str, Any], output: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"
    ws["A1"] = "CONCILIACIÓN AUTOMÁTICA DE CRÉDITOS"
    ws["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = HEADER_FILL
    ws.merge_cells("A1:D1")
    ws["A2"] = "Ejecución"
    ws["B2"] = datetime.now()
    ws["B2"].number_format = "dd/mm/yyyy hh:mm"
    ws["A4"] = "Indicador"; ws["B4"] = "Resultado"
    for c in ws[4]: c.fill, c.font = HEADER_FILL, WHITE_FONT
    row = 5
    for label, value in result["metrics"].items():
        ws.cell(row, 1, label); ws.cell(row, 2, value)
        if "Saldo" in label: ws.cell(row, 2).number_format = "#,##0.00"
        row += 1
    ws["A14"] = "Interpretación"
    ws["A14"].fill, ws["A14"].font = SUBHEADER_FILL, Font(bold=True)
    notes = [
        "NUEVO EN SIGNO: agregar o validar en la cartera.",
        "VALIDAR PAGO/ANULACIÓN: estaba en Créditos actuales, pero ya no aparece en el reporte pendiente de SIGNO.",
        "La desaparición de SIGNO no confirma por sí sola el pago; requiere validación humana.",
        "El proceso nunca modifica los dos archivos originales.",
    ]
    for i, note in enumerate(notes, start=15): ws.cell(i, 1, f"• {note}")
    ws.column_dimensions["A"].width = 56; ws.column_dimensions["B"].width = 24

    master_cols = [("ID","ID"),("RUC","RUC"),("CLIENTE","Cliente"),("TIPO","Tipo"),("SERIE","Serie"),("DOC","Documento"),("EMISION","Emisión"),("TRAMITE","Trámite"),("FACTURADO_SIGNO","Facturado SIGNO"),("SALDO_SIGNO","Saldo SIGNO"),("SALDO_MANUAL","Saldo manual"),("VENCIMIENTO","Vencimiento"),("ESTADO_CONCILIACION","Estado conciliación"),("PRIORIDAD","Prioridad"),("DIAS_PARA_VENCER","Días para vencer"),("DETALLE_COBRANZA","Detalle cobranza"),("SOLICITANTE","Solicitante"),("ULTIMA_SINCRONIZACION","Última sincronización")]
    signo_cols = [("ID","ID"),("RUC","RUC"),("CLIENTE","Cliente"),("TIPO","Tipo"),("SERIE","Serie"),("DOC","Documento"),("EMISION","Emisión"),("TRAMITE","Trámite"),("FACTURADO","Facturado"),("SALDO","Saldo"),("VENCIMIENTO","Vencimiento")]
    old_cols = [("ID","ID"),("RUC","RUC"),("CLIENTE","Cliente"),("TIPO","Tipo"),("SERIE","Serie"),("DOC","Documento"),("EMISION","Emisión"),("TRAMITE","Trámite"),("SALDO_MANUAL","Saldo manual"),("VENCIMIENTO","Vencimiento"),("DETALLE_COBRANZA","Detalle cobranza"),("SOLICITANTE","Solicitante")]
    change_cols = [("ID","ID"),("CLIENTE","Cliente"),("SALDO_ANTERIOR","Saldo anterior"),("SALDO_ACTUAL","Saldo actual"),("DIFERENCIA","Diferencia")]
    collection_cols = [("PRIORIDAD","Prioridad"),("CLIENTE","Cliente"),("RUC","RUC"),("TIPO","Tipo"),("SERIE","Serie"),("DOC","Documento"),("SALDO_SIGNO","Saldo"),("VENCIMIENTO","Vencimiento"),("DIAS_PARA_VENCER","Días para vencer"),("TRAMITE","Trámite"),("SOLICITANTE","Solicitante"),("DETALLE_COBRANZA","Detalle cobranza")]
    add_table_sheet(wb,"Base_Maestra",result["master"],master_cols,"tblBaseMaestra")
    add_table_sheet(wb,"Nuevos_SIGNO",result["new"],signo_cols,"tblNuevos")
    add_table_sheet(wb,"Posibles_Pagos",result["possible_paid"],old_cols,"tblPosiblesPagos")
    add_table_sheet(wb,"Saldos_Modificados",result["changes"],change_cols,"tblCambiosSaldo")
    add_table_sheet(wb,"Cobranza_Prioritaria",result["collections"],collection_cols,"tblCobranza")
    if result["duplicates"]:
        add_table_sheet(wb,"Duplicados",[{"DETALLE": x} for x in result["duplicates"]],[("DETALLE","Registro duplicado")],"tblDuplicados")

    wb.save(output) 


def choose_path(explicit: Optional[str], default_name: str, base: Path) -> Path:
    if explicit:
        return Path(explicit).expanduser().resolve()
    exact = base / default_name
    if exact.exists(): return exact
    candidates = [p for p in base.glob("*.xlsx") if not p.name.startswith("~$") and "MVP_" not in p.name and "CONCILIACION_" not in p.name]
    keyword = "SIGNO" if "SIGNO" in default_name else "ACTUALES"
    matches = [p for p in candidates if keyword in normalize_text(p.name)]
    if len(matches) == 1: return matches[0]
    return exact


def main() -> int:
    parser = argparse.ArgumentParser(description="Compara el reporte pendiente de SIGNO con Créditos actuales.")
    parser.add_argument("--signo", help="Ruta del reporte exportado de SIGNO")
    parser.add_argument("--creditos", help="Ruta del Excel Créditos actuales")
    parser.add_argument("--salida", help="Ruta del archivo de salida")
    parser.add_argument("--sin-pausa", action="store_true", help="No esperar Enter al finalizar")
    args = parser.parse_args()

    base = Path(__file__).resolve().parent
    signo_path = choose_path(args.signo, DEFAULT_SIGNO, base)
    creditos_path = choose_path(args.creditos, DEFAULT_CREDITOS, base)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output = Path(args.salida).expanduser().resolve() if args.salida else base / f"CONCILIACION_CREDITOS_COMPATIBLE_{timestamp}.xlsx"
    
    try:
        if not signo_path.exists(): raise FileNotFoundError(f"No se encontró el reporte de SIGNO: {signo_path}")
        if not creditos_path.exists(): raise FileNotFoundError(f"No se encontró Créditos actuales: {creditos_path}")
        print("Leyendo reporte de SIGNO...")
        signo = read_signo(signo_path)
        print("Leyendo Créditos actuales...")
        creditos = read_creditos(creditos_path)
        print("Conciliando comprobantes...")
        result = reconcile(signo, creditos)
        write_output(result, output)
        print("\nPROCESO TERMINADO")
        for label, value in result["metrics"].items():
            if isinstance(value, float): print(f"- {label}: S/ {value:,.2f}")
            else: print(f"- {label}: {value}")
        print(f"\nArchivo generado:\n{output}")
        return 0
    except Exception as exc:
        print(f"\nERROR: {exc}")
        print("Revise que los Excel estén cerrados, tengan los nombres correctos y conserven sus encabezados.")
        return 1
    finally:
        if not args.sin_pausa and sys.stdin.isatty():
            input("\nPresione Enter para cerrar...")


if __name__ == "__main__":
    raise SystemExit(main())