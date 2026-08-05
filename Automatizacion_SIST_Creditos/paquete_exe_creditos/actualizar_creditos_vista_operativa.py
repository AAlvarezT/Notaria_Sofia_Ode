#!/usr/bin/env python3
"""
Actualización de Créditos Actuales desde el reporte de SIGNO.

OBJETIVO
- Mantener una vista operativa parecida a "CREDITOS ACTUALES 2026.xlsx".
- Agrupar por responsable/sección y por cliente.
- Conservar comentarios, solicitante y marcas manuales.
- Añadir nuevos comprobantes, actualizar saldos y señalar posibles pagos.
- Generar subtotales, resumen, revisión de cambios, historial y recordatorios.

REQUISITO
    pip install openpyxl

USO
    python actualizar_creditos_vista_operativa.py

O con rutas explícitas:
    python actualizar_creditos_vista_operativa.py \
      --signo "CREDITOS SISTEMA SIGNO.xlsx" \
      --creditos "CREDITOS ACTUALES 2026.xlsx" \
      --salida "CREDITOS ACTUALIZADOS.xlsx"

El programa NO sobrescribe los archivos originales.
"""

from __future__ import annotations

import argparse
import re
import shutil
import sys
import unicodedata
from collections import Counter, defaultdict
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: falta openpyxl. Instálelo con: pip install openpyxl")
    input("Presione Enter para cerrar...")
    raise SystemExit(1)

DEFAULT_SIGNO = "CREDITOS SISTEMA SIGNO.xlsx"
DEFAULT_CREDITOS = "CREDITOS ACTUALES 2026.xlsx"

# Colores simples y compatibles con Excel de escritorio.
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
SECTION_BLUE = "DDEBF7"
CLIENT_BLUE = "EAF2F8"
TOTAL_BLUE = "BDD7EE"
YELLOW = "FFF2CC"
RED = "F4CCCC"
GREEN = "D9EAD3"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
BLACK = "000000"
THIN = Side(style="thin", color="B7C9D6")

HEADERS = [
    "", "RUC/DNI", "RAZON SOCIAL", "EMISION", "TIPO", "SERIE", "DOC",
    "TRAMITE", "PENDIENTE DE PAGO NOTARIAL", "PENDIENTE DE PAGO REGISTRAL",
    "PENDIENTE DE PAGO GESTION", "VENCIMIENTO", "DETALLE DE COBRANZA",
    "DIAS POR VENCER", "SOLICITANTE", "", "ARCHIVO"
]


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return re.sub(r"\s+", " ", text)


def normalize_text(value: Any) -> str:
    text = clean_text(value).upper()
    text = unicodedata.normalize("NFKD", text)
    return "".join(c for c in text if not unicodedata.combining(c))


def normalize_ruc(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\D", "", str(value))


def normalize_doc(value: Any, width: int = 7) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    raw = clean_text(value)
    digits = re.sub(r"\D", "", raw)
    return digits.zfill(width) if digits else normalize_text(raw)


def parse_date(value: Any) -> Optional[datetime]:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    # Excel serial simple, por si data_only devuelve número.
    if isinstance(value, (int, float)) and value > 1000:
        try:
            return datetime(1899, 12, 30) + __import__('datetime').timedelta(days=float(value))
        except Exception:
            pass
    text = clean_text(value)
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def to_number(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = clean_text(value).replace("S/", "").replace(",", "")
    try:
        return float(text)
    except ValueError:
        return 0.0


def make_key(tipo: Any, serie: Any, documento: Any) -> str:
    t = normalize_text(tipo)
    s = normalize_text(serie)
    d = normalize_doc(documento)
    return f"{t}|{s}|{d}" if t and s and d else ""


def client_key(ruc: Any, cliente: Any) -> str:
    nruc = normalize_ruc(ruc)
    return f"RUC:{nruc}" if nruc else f"NOMBRE:{normalize_text(cliente)}"


def find_header(ws, aliases: Dict[str, str], max_rows: int = 25, max_cols: int = 30):
    for r in range(1, min(ws.max_row, max_rows) + 1):
        found: Dict[str, int] = {}
        for c in range(1, min(ws.max_column, max_cols) + 1):
            n = normalize_text(ws.cell(r, c).value)
            if n in aliases:
                found[aliases[n]] = c
        if all(k in found for k in ("TIPO", "SERIE", "DOC")):
            return r, found
    raise ValueError(f"No se encontró la cabecera esperada en la hoja {ws.title}.")


def read_signo(path: Path) -> List[Dict[str, Any]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    aliases = {
        "RUC": "RUC", "GRUPO": "GRUPO", "EMISION": "EMISION", "TIPO": "TIPO",
        "SERIE A": "SERIE", "SERIE": "SERIE", "DOC": "DOC", "CLIENTE": "CLIENTE",
        "KARDEX": "TRAMITE", "TRAMITE": "TRAMITE", "FACTURADO": "FACTURADO",
        "SALDO": "SALDO", "VENCIMIENTO": "VENCIMIENTO", "EMPLEADO": "EMPLEADO",
        "COMENTARIOS": "COMENTARIO_SIGNO"
    }
    header, cols = find_header(ws, aliases)
    out = []
    for r in range(header + 1, ws.max_row + 1):
        key = make_key(ws.cell(r, cols["TIPO"]).value, ws.cell(r, cols["SERIE"]).value, ws.cell(r, cols["DOC"]).value)
        if not key:
            continue
        cliente = clean_text(ws.cell(r, cols.get("CLIENTE", 0)).value) if cols.get("CLIENTE") else ""
        grupo = clean_text(ws.cell(r, cols.get("GRUPO", 0)).value) if cols.get("GRUPO") else ""
        out.append({
            "ID": key,
            "RUC": normalize_ruc(ws.cell(r, cols.get("RUC", 0)).value) if cols.get("RUC") else "",
            "CLIENTE": cliente or grupo,
            "GRUPO_SIGNO": grupo,
            "EMISION": parse_date(ws.cell(r, cols.get("EMISION", 0)).value) if cols.get("EMISION") else None,
            "TIPO": normalize_text(ws.cell(r, cols["TIPO"]).value),
            "SERIE": normalize_text(ws.cell(r, cols["SERIE"]).value),
            "DOC": normalize_doc(ws.cell(r, cols["DOC"]).value),
            "TRAMITE": clean_text(ws.cell(r, cols.get("TRAMITE", 0)).value) if cols.get("TRAMITE") else "",
            "FACTURADO": to_number(ws.cell(r, cols.get("FACTURADO", 0)).value) if cols.get("FACTURADO") else 0,
            "SALDO": to_number(ws.cell(r, cols.get("SALDO", 0)).value) if cols.get("SALDO") else 0,
            "VENCIMIENTO": parse_date(ws.cell(r, cols.get("VENCIMIENTO", 0)).value) if cols.get("VENCIMIENTO") else None,
            "EMPLEADO": clean_text(ws.cell(r, cols.get("EMPLEADO", 0)).value) if cols.get("EMPLEADO") else "",
            "COMENTARIO_SIGNO": clean_text(ws.cell(r, cols.get("COMENTARIO_SIGNO", 0)).value) if cols.get("COMENTARIO_SIGNO") else "",
        })
    wb.close()
    return out


def read_current_view(path: Path) -> Tuple[List[Dict[str, Any]], List[str], Dict[str, int], Dict[str, int]]:
    """Lee la vista actual, incluyendo la sección/responsable que precede cada bloque."""
    wb = load_workbook(path, data_only=False, read_only=False)
    ws = wb["CREDITOS"] if "CREDITOS" in wb.sheetnames else wb.active
    aliases = {
        "RUC/DNI": "RUC", "RUC": "RUC", "RAZON SOCIAL": "CLIENTE", "RAZON SOCIAL ": "CLIENTE",
        "EMISION": "EMISION", "TIPO": "TIPO", "SERIE": "SERIE", "DOC": "DOC",
        "TRAMITE": "TRAMITE", "PENDIENTE DE PAGO NOTARIAL": "NOTARIAL",
        "PENDIENTE DE PAGO REGISTRAL": "REGISTRAL", "PENDIENTE DE PAGO GESTION": "GESTION",
        "VENCIMIENTO": "VENCIMIENTO", "DETALLE DE COBRANZA": "DETALLE",
        "DIAS POR VENCER": "DIAS", "SOLICITANTE": "SOLICITANTE", "ARCHIVO": "ARCHIVO"
    }
    header, cols = find_header(ws, aliases)
    records: List[Dict[str, Any]] = []
    section_order: List[str] = []
    client_order: Dict[str, int] = {}
    section_rank: Dict[str, int] = {}
    current_section = "SIN RESPONSABLE"
    seen_clients = 0

    for r in range(header + 1, ws.max_row + 1):
        tipo = ws.cell(r, cols["TIPO"]).value
        serie = ws.cell(r, cols["SERIE"]).value
        doc = ws.cell(r, cols["DOC"]).value
        key = make_key(tipo, serie, doc)

        # Las filas de sección tienen texto en Razón Social, sin documento ni montos.
        cliente_cell = clean_text(ws.cell(r, cols.get("CLIENTE", 3)).value)
        has_amount = any(to_number(ws.cell(r, cols.get(k, 0)).value) for k in ("NOTARIAL", "REGISTRAL", "GESTION") if cols.get(k))
        if not key and cliente_cell and not has_amount:
            current_section = cliente_cell
            if current_section not in section_rank:
                section_rank[current_section] = len(section_order)
                section_order.append(current_section)
            continue
        if not key:
            continue

        notarial = to_number(ws.cell(r, cols.get("NOTARIAL", 0)).value) if cols.get("NOTARIAL") else 0
        registral = to_number(ws.cell(r, cols.get("REGISTRAL", 0)).value) if cols.get("REGISTRAL") else 0
        gestion = to_number(ws.cell(r, cols.get("GESTION", 0)).value) if cols.get("GESTION") else 0
        ruc = normalize_ruc(ws.cell(r, cols.get("RUC", 0)).value) if cols.get("RUC") else ""
        ck = client_key(ruc, cliente_cell)
        if ck not in client_order:
            client_order[ck] = seen_clients
            seen_clients += 1
        records.append({
            "ID": key, "RUC": ruc, "CLIENTE": cliente_cell,
            "EMISION": parse_date(ws.cell(r, cols.get("EMISION", 0)).value) if cols.get("EMISION") else None,
            "TIPO": normalize_text(tipo), "SERIE": normalize_text(serie), "DOC": normalize_doc(doc),
            "TRAMITE": clean_text(ws.cell(r, cols.get("TRAMITE", 0)).value) if cols.get("TRAMITE") else "",
            "NOTARIAL": notarial, "REGISTRAL": registral, "GESTION": gestion,
            "SALDO_MANUAL": notarial + registral + gestion,
            "VENCIMIENTO": parse_date(ws.cell(r, cols.get("VENCIMIENTO", 0)).value) if cols.get("VENCIMIENTO") else None,
            "DETALLE_COBRANZA": clean_text(ws.cell(r, cols.get("DETALLE", 0)).value) if cols.get("DETALLE") else "",
            "DIAS_MANUAL": clean_text(ws.cell(r, cols.get("DIAS", 0)).value) if cols.get("DIAS") else "",
            "SOLICITANTE": clean_text(ws.cell(r, cols.get("SOLICITANTE", 0)).value) if cols.get("SOLICITANTE") else "",
            "MARCA": clean_text(ws.cell(r, 16).value),
            "ARCHIVO": clean_text(ws.cell(r, cols.get("ARCHIVO", 17)).value) if cols.get("ARCHIVO") else "",
            "RESPONSABLE": current_section,
            "ORDEN_ORIGINAL": r,
            "CLIENT_KEY": ck,
        })
    wb.close()
    return records, section_order, section_rank, client_order


def majority_category(records: Iterable[Dict[str, Any]]) -> str:
    counts = Counter()
    for x in records:
        values = {"NOTARIAL": x.get("NOTARIAL", 0), "REGISTRAL": x.get("REGISTRAL", 0), "GESTION": x.get("GESTION", 0)}
        cat = max(values, key=values.get)
        if values[cat] > 0:
            counts[cat] += 1
    return counts.most_common(1)[0][0] if counts else "NOTARIAL"


def reconcile(signo_rows: List[Dict[str, Any]], old_rows: List[Dict[str, Any]], section_order: List[str], client_order: Dict[str, int]):
    signo = {x["ID"]: x for x in signo_rows}
    old = {x["ID"]: x for x in old_rows}

    by_client: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for x in old_rows:
        by_client[x["CLIENT_KEY"]].append(x)

    # Índices secundarios para asignar clientes nuevos aun si falta el RUC en el archivo manual.
    client_name_index: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    ruc_index: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for x in old_rows:
        client_name_index[normalize_text(x["CLIENTE"])].append(x)
        if x["RUC"]:
            ruc_index[x["RUC"]].append(x)

    merged = []
    changes = []
    exceptions = []
    next_order = max(client_order.values(), default=-1) + 1

    for key in sorted(set(signo) | set(old)):
        s, o = signo.get(key), old.get(key)
        if s and o:
            row = dict(o)
            previous = o["SALDO_MANUAL"]
            current = s["SALDO"]
            row.update({
                "RUC": s["RUC"] or o["RUC"],
                "CLIENTE": s["CLIENTE"] or o["CLIENTE"],
                "EMISION": s["EMISION"] or o["EMISION"],
                "TRAMITE": s["TRAMITE"] or o["TRAMITE"],
                "VENCIMIENTO": s["VENCIMIENTO"] or o["VENCIMIENTO"],
                "SALDO_SIGNO": current,
                "ESTADO": "PENDIENTE",
                "COMENTARIO_SIGNO": s["COMENTARIO_SIGNO"],
            })
            # Mantiene la categoría manual, actualizando el monto total a saldo SIGNO.
            cat = majority_category([o])
            row["NOTARIAL"] = current if cat == "NOTARIAL" else 0
            row["REGISTRAL"] = current if cat == "REGISTRAL" else 0
            row["GESTION"] = current if cat == "GESTION" else 0
            if abs(previous - current) > 0.009:
                changes.append(["SALDO MODIFICADO", key, row["RUC"], row["CLIENTE"], previous, current, "Revisar diferencia"])
            merged.append(row)

        elif s and not o:
            candidates = ruc_index.get(s["RUC"], []) if s["RUC"] else []
            if not candidates:
                candidates = client_name_index.get(normalize_text(s["CLIENTE"]), [])
            if candidates:
                base = candidates[0]
                responsable = base["RESPONSABLE"]
                ck = base["CLIENT_KEY"]
                order = client_order.get(ck, next_order)
                cat = majority_category(candidates)
            else:
                responsable = "NUEVOS / SIN ASIGNAR"
                ck = client_key(s["RUC"], s["CLIENTE"])
                if ck not in client_order:
                    client_order[ck] = next_order
                    next_order += 1
                order = client_order[ck]
                cat = "NOTARIAL"
                exceptions.append(["CLIENTE SIN ASIGNACIÓN", key, s["RUC"], s["CLIENTE"], "Asignado temporalmente a NUEVOS / SIN ASIGNAR y saldo en Notarial"])
            row = {
                **s,
                "RESPONSABLE": responsable, "CLIENT_KEY": ck, "ORDEN_ORIGINAL": 999999,
                "NOTARIAL": s["SALDO"] if cat == "NOTARIAL" else 0,
                "REGISTRAL": s["SALDO"] if cat == "REGISTRAL" else 0,
                "GESTION": s["SALDO"] if cat == "GESTION" else 0,
                "SALDO_MANUAL": 0, "SALDO_SIGNO": s["SALDO"], "DETALLE_COBRANZA": "",
                "DIAS_MANUAL": "", "SOLICITANTE": s["EMPLEADO"], "MARCA": "", "ARCHIVO": "",
                "ESTADO": "NUEVO", "CLIENT_ORDER": order,
            }
            merged.append(row)
            changes.append(["NUEVO", key, s["RUC"], s["CLIENTE"], 0, s["SALDO"], "Ingresado automáticamente"])

        elif o and not s:
            row = dict(o)
            row.update({"SALDO_SIGNO": 0, "ESTADO": "VALIDAR PAGO/ANULACION", "COMENTARIO_SIGNO": ""})
            # Se mantiene visible y con su monto anterior hasta validación humana.
            merged.append(row)
            changes.append(["POSIBLE PAGO/ANULACIÓN", key, o["RUC"], o["CLIENTE"], o["SALDO_MANUAL"], 0, "Ya no aparece en el reporte SIGNO"])

    # Inserta sección nueva al final.
    if any(x["RESPONSABLE"] == "NUEVOS / SIN ASIGNAR" for x in merged) and "NUEVOS / SIN ASIGNAR" not in section_order:
        section_order.append("NUEVOS / SIN ASIGNAR")

    return merged, changes, exceptions


def clear_sheet(ws):
    for row in ws.iter_rows():
        for cell in row:
            cell.value = None
            cell._style = copy(cell._style)
    if ws.merged_cells.ranges:
        for rng in list(ws.merged_cells.ranges):
            ws.unmerge_cells(str(rng))


def setup_columns(ws):
    widths = {"A": 3, "B": 16, "C": 38, "D": 12, "E": 8, "F": 10, "G": 13, "H": 23,
              "I": 18, "J": 18, "K": 18, "L": 13, "M": 45, "N": 14, "O": 22, "P": 5, "Q": 12}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "B5"
    ws.sheet_view.showGridLines = False


def style_header(ws, row):
    for c in range(2, 18):
        cell = ws.cell(row, c)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color=WHITE, bold=True, size=9)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    ws.row_dimensions[row].height = 34


def write_operational_view(ws, records, section_order, client_order):
    clear_sheet(ws)
    setup_columns(ws)
    ws.merge_cells("B1:Q1")
    ws["B1"] = "NOTARIA SOFIA INES ODE PEREYRA y TSM"
    ws["B1"].font = Font(bold=True, size=14, color=BLUE)
    ws["B1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("B2:Q2")
    ws["B2"] = "Reporte de Comprobantes Pendientes de Cancelar"
    ws["B2"].font = Font(bold=True, size=12)
    ws["B2"].alignment = Alignment(horizontal="center")
    ws["C3"] = datetime.now()
    ws["C3"].number_format = "dd/mm/yyyy hh:mm"
    for c, val in enumerate(HEADERS, 1):
        ws.cell(4, c).value = val
    style_header(ws, 4)

    section_rank = {name: i for i, name in enumerate(section_order)}
    grouped: Dict[str, Dict[str, List[Dict[str, Any]]]] = defaultdict(lambda: defaultdict(list))
    for rec in records:
        grouped[rec.get("RESPONSABLE") or "SIN RESPONSABLE"][rec["CLIENT_KEY"]].append(rec)

    row = 5
    for section in sorted(grouped, key=lambda x: section_rank.get(x, 9999)):
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=17)
        ws.cell(row, 3).value = section
        ws.cell(row, 3).font = Font(bold=True, size=11, color=BLUE)
        ws.cell(row, 3).fill = PatternFill("solid", fgColor=SECTION_BLUE)
        ws.cell(row, 3).alignment = Alignment(horizontal="left")
        row += 1

        clients = grouped[section]
        sorted_clients = sorted(clients.items(), key=lambda kv: client_order.get(kv[0], 999999))
        for ck, items in sorted_clients:
            items.sort(key=lambda x: (x.get("EMISION") or datetime.max, x["ID"]))
            # Un encabezado ligero por cliente mejora la lectura sin romper el formato familiar.
            ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=17)
            ws.cell(row, 3).value = items[0].get("CLIENTE") or "CLIENTE SIN NOMBRE"
            ws.cell(row, 3).font = Font(bold=True, size=10)
            ws.cell(row, 3).fill = PatternFill("solid", fgColor=CLIENT_BLUE)
            row += 1
            start_detail = row
            for rec in items:
                venc = rec.get("VENCIMIENTO")
                days = (venc.date() - date.today()).days if venc else None
                values = [
                    "", rec.get("RUC", ""), rec.get("CLIENTE", ""), rec.get("EMISION"), rec.get("TIPO", ""),
                    rec.get("SERIE", ""), rec.get("DOC", ""), rec.get("TRAMITE", ""), rec.get("NOTARIAL", 0),
                    rec.get("REGISTRAL", 0), rec.get("GESTION", 0), venc, rec.get("DETALLE_COBRANZA", ""),
                    days if days is not None else rec.get("DIAS_MANUAL", ""), rec.get("SOLICITANTE", ""),
                    rec.get("MARCA", ""), rec.get("ARCHIVO", "")
                ]
                for c, value in enumerate(values, 1):
                    ws.cell(row, c).value = value
                for c in range(2, 18):
                    ws.cell(row, c).border = Border(bottom=Side(style="hair", color="D9E1F2"))
                    ws.cell(row, c).alignment = Alignment(vertical="top", wrap_text=(c in (3, 8, 13, 15)))
                for c in (9, 10, 11):
                    ws.cell(row, c).number_format = '#,##0.00'
                for c in (4, 12):
                    ws.cell(row, c).number_format = 'dd/mm/yyyy'
                ws.cell(row, 7).number_format = '@'
                status = rec.get("ESTADO", "")
                if status == "NUEVO":
                    fill = YELLOW
                elif status == "VALIDAR PAGO/ANULACION":
                    fill = GREEN
                    ws.cell(row, 13).value = (rec.get("DETALLE_COBRANZA", "") + " | VALIDAR PAGO/ANULACIÓN").strip(" |")
                else:
                    fill = None
                if fill:
                    for c in range(2, 18):
                        ws.cell(row, c).fill = PatternFill("solid", fgColor=fill)
                row += 1

            # Subtotal por cliente.
            ws.cell(row, 3).value = f"SUBTOTAL {items[0].get('CLIENTE','')}"
            ws.cell(row, 3).font = Font(bold=True)
            for c in (9, 10, 11):
                letter = get_column_letter(c)
                ws.cell(row, c).value = f"=SUM({letter}{start_detail}:{letter}{row-1})"
                ws.cell(row, c).number_format = '#,##0.00'
                ws.cell(row, c).font = Font(bold=True)
            for c in range(2, 18):
                ws.cell(row, c).fill = PatternFill("solid", fgColor=TOTAL_BLUE)
                ws.cell(row, c).border = Border(top=THIN, bottom=THIN)
            row += 1
        row += 1
    ws.auto_filter.ref = f"B4:Q{max(4,row-1)}"


def replace_or_create(wb, name):
    if name in wb.sheetnames:
        ws = wb[name]
        # remove all content/merges safely by replacing the sheet, preserving order approximately.
        idx = wb.sheetnames.index(name)
        wb.remove(ws)
        ws = wb.create_sheet(name, idx)
    else:
        ws = wb.create_sheet(name)
    return ws


def write_review(ws, changes, exceptions):
    headers = ["TIPO DE CAMBIO", "ID", "RUC", "CLIENTE", "SALDO ANTERIOR", "SALDO SIGNO", "ACCIÓN / OBSERVACIÓN"]
    ws.append(headers)
    for row in changes + exceptions:
        # Completar filas de excepción al ancho esperado.
        ws.append(row + [""] * (len(headers) - len(row)))
    for c in range(1, len(headers)+1):
        ws.cell(1,c).fill = PatternFill("solid", fgColor=BLUE)
        ws.cell(1,c).font = Font(color=WHITE,bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{max(1,ws.max_row)}"
    for col, width in zip("ABCDEFG", [24,24,16,38,18,18,45]):
        ws.column_dimensions[col].width = width
    for row in range(2, ws.max_row+1):
        ws.cell(row,5).number_format = '#,##0.00'
        ws.cell(row,6).number_format = '#,##0.00'


def write_base(ws, records):
    headers = ["ID","RUC","CLIENTE","RESPONSABLE","EMISION","TIPO","SERIE","DOC","TRAMITE","NOTARIAL","REGISTRAL","GESTION","SALDO_SIGNO","VENCIMIENTO","DETALLE_COBRANZA","SOLICITANTE","ARCHIVO","ESTADO","ULTIMA_ACTUALIZACION"]
    ws.append(headers)
    for x in records:
        ws.append([x.get(k, "") for k in headers[:-1]] + [datetime.now()])
    for c in range(1,len(headers)+1):
        ws.cell(1,c).fill=PatternFill("solid",fgColor=BLUE); ws.cell(1,c).font=Font(color=WHITE,bold=True)
    ws.freeze_panes="A2"
    ws.auto_filter.ref=f"A1:S{max(1,ws.max_row)}"
    for c in (5,14,19):
        for r in range(2,ws.max_row+1): ws.cell(r,c).number_format='dd/mm/yyyy'
    for c in (10,11,12,13):
        for r in range(2,ws.max_row+1): ws.cell(r,c).number_format='#,##0.00'
    ws.sheet_state = "hidden"


def write_summary(ws, records):
    agg = defaultdict(lambda: [0.0,0.0,0.0,0, None, ""])
    for x in records:
        k=(x.get("RESPONSABLE",""), x.get("CLIENT_KEY",""), x.get("CLIENTE",""), x.get("RUC",""))
        agg[k][0]+=x.get("NOTARIAL",0); agg[k][1]+=x.get("REGISTRAL",0); agg[k][2]+=x.get("GESTION",0); agg[k][3]+=1
        v=x.get("VENCIMIENTO")
        if v and (agg[k][4] is None or v<agg[k][4]): agg[k][4]=v
    headers=["RESPONSABLE","RUC","CLIENTE","NOTARIAL","REGISTRAL","GESTION","TOTAL","N° COMPROBANTES","VENCIMIENTO MÁS ANTIGUO"]
    ws.append(headers)
    for (resp,ck,cli,ruc),v in sorted(agg.items(), key=lambda z:(z[0][0],z[0][2])):
        ws.append([resp,ruc,cli,v[0],v[1],v[2],v[0]+v[1]+v[2],v[3],v[4]])
    for c in range(1,len(headers)+1): ws.cell(1,c).fill=PatternFill("solid",fgColor=BLUE); ws.cell(1,c).font=Font(color=WHITE,bold=True)
    ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:I{max(1,ws.max_row)}"
    for c in (4,5,6,7):
        for r in range(2,ws.max_row+1): ws.cell(r,c).number_format='#,##0.00'
    for r in range(2,ws.max_row+1): ws.cell(r,9).number_format='dd/mm/yyyy'
    widths=[28,16,40,15,15,15,15,16,20]
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w


def write_reminders(ws, records):
    headers=["PRIORIDAD","RESPONSABLE","RUC","CLIENTE","COMPROBANTE","SALDO","VENCIMIENTO","DÍAS","DETALLE ACTUAL","PRÓXIMA GESTIÓN","MENSAJE SUGERIDO"]
    ws.append(headers)
    today=date.today()
    rows=[]
    for x in records:
        if x.get("ESTADO") == "VALIDAR PAGO/ANULACION":
            priority="VALIDAR"
        elif not x.get("VENCIMIENTO"):
            priority="SIN FECHA"
        else:
            d=(x["VENCIMIENTO"].date()-today).days
            if d < -30: priority="CRÍTICA"
            elif d < 0: priority="VENCIDA"
            elif d <= 3: priority="ALTA"
            elif d <= 7: priority="PREVENTIVA"
            else: continue
        saldo=x.get("NOTARIAL",0)+x.get("REGISTRAL",0)+x.get("GESTION",0)
        dias=(x["VENCIMIENTO"].date()-today).days if x.get("VENCIMIENTO") else ""
        comprobante=f"{x.get('TIPO','')}-{x.get('SERIE','')}-{x.get('DOC','')}"
        mensaje=f"Estimados señores de {x.get('CLIENTE','')}, les recordamos el comprobante {comprobante} por S/ {saldo:,.2f}."
        rows.append([priority,x.get("RESPONSABLE",""),x.get("RUC",""),x.get("CLIENTE",""),comprobante,saldo,x.get("VENCIMIENTO"),dias,x.get("DETALLE_COBRANZA",""),"",mensaje])
    rank={"VALIDAR":0,"CRÍTICA":1,"VENCIDA":2,"ALTA":3,"PREVENTIVA":4,"SIN FECHA":5}
    rows.sort(key=lambda z:(rank.get(z[0],99), z[7] if isinstance(z[7],int) else 9999))
    for r in rows: ws.append(r)
    for c in range(1,len(headers)+1): ws.cell(1,c).fill=PatternFill("solid",fgColor=BLUE); ws.cell(1,c).font=Font(color=WHITE,bold=True)
    ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:K{max(1,ws.max_row)}"
    for r in range(2,ws.max_row+1):
        ws.cell(r,6).number_format='#,##0.00'; ws.cell(r,7).number_format='dd/mm/yyyy'
    widths=[14,28,16,38,24,14,14,10,42,18,70]
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w


def ensure_history(wb):
    if "HISTORIAL_GESTIONES" not in wb.sheetnames:
        ws=wb.create_sheet("HISTORIAL_GESTIONES")
        ws.append(["FECHA","ID COMPROBANTE","RUC","CLIENTE","TIPO DE GESTIÓN","DETALLE","USUARIO","PRÓXIMA GESTIÓN"])
        for c in range(1,9): ws.cell(1,c).fill=PatternFill("solid",fgColor=BLUE); ws.cell(1,c).font=Font(color=WHITE,bold=True)
        ws.freeze_panes="A2"


def main():
    parser=argparse.ArgumentParser(description="Actualiza Créditos actuales desde SIGNO conservando vista operativa.")
    parser.add_argument("--signo", default=DEFAULT_SIGNO)
    parser.add_argument("--creditos", default=DEFAULT_CREDITOS)
    parser.add_argument("--salida", default="")
    args=parser.parse_args()
    signo_path=Path(args.signo).expanduser().resolve()
    creditos_path=Path(args.creditos).expanduser().resolve()
    if not signo_path.exists(): raise FileNotFoundError(f"No existe: {signo_path}")
    if not creditos_path.exists(): raise FileNotFoundError(f"No existe: {creditos_path}")
    output=Path(args.salida).expanduser().resolve() if args.salida else creditos_path.with_name(f"CREDITOS_ACTUALIZADOS_{datetime.now():%Y%m%d_%H%M%S}.xlsx")

    print("1/5 Leyendo reporte SIGNO...")
    signo=read_signo(signo_path)
    print("2/5 Leyendo vista Créditos actuales y datos manuales...")
    old, section_order, section_rank, client_order=read_current_view(creditos_path)
    print("3/5 Conciliando comprobantes...")
    merged, changes, exceptions=reconcile(signo, old, section_order, client_order)

    print("4/5 Generando vista operativa y hojas de control...")
    shutil.copy2(creditos_path, output)
    wb=load_workbook(output)
    ws_creditos=replace_or_create(wb,"CREDITOS")
    write_operational_view(ws_creditos, merged, section_order, client_order)
    write_summary(replace_or_create(wb,"RESUMEN"), merged)
    write_review(replace_or_create(wb,"REVISAR_CAMBIOS"), changes, exceptions)
    write_base(replace_or_create(wb,"BASE_DATOS"), merged)
    write_reminders(replace_or_create(wb,"RECORDATORIOS"), merged)
    ensure_history(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(output)
    wb.close()

    nuevos=sum(1 for x in changes if x[0]=="NUEVO")
    pagos=sum(1 for x in changes if x[0]=="POSIBLE PAGO/ANULACIÓN")
    saldos=sum(1 for x in changes if x[0]=="SALDO MODIFICADO")
    print("5/5 Proceso terminado.")
    print(f"Salida: {output}")
    print(f"Comprobantes SIGNO: {len(signo)}")
    print(f"Registros previos: {len(old)}")
    print(f"Nuevos: {nuevos} | Saldos modificados: {saldos} | Posibles pagos/anulaciones: {pagos}")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"\nERROR: {exc}")
        import traceback
        traceback.print_exc()
        input("Presione Enter para cerrar...")
        sys.exit(1)
