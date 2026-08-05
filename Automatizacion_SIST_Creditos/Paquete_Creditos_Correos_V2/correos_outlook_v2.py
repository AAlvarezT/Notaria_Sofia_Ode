from __future__ import annotations

import html
import re
import csv
import os
import winreg
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Tuple

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

BLUE = "1F4E78"
WHITE = "FFFFFF"


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _norm(value: Any) -> str:
    text = _text(value).upper()
    text = re.sub(r"\s+", " ", text)
    return text


def _ruc(value: Any) -> str:
    digits = re.sub(r"\D", "", _text(value))
    return digits


def _yes(value: Any, default: bool = True) -> bool:
    text = _norm(value)
    if not text:
        return default
    return text in {"SI", "SÍ", "YES", "TRUE", "1", "X", "ACTIVO"}


def _valid_email(value: str) -> bool:
    return bool(re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", value.strip()))


def _split_emails(value: Any) -> List[str]:
    raw = _text(value)
    if not raw:
        return []
    parts = re.split(r"[;,\s]+", raw)
    return [p for p in parts if p and _valid_email(p)]


def ensure_email_sheets(wb) -> None:
    if "MAESTRO_CLIENTES" not in wb.sheetnames:
        ws = wb.create_sheet("MAESTRO_CLIENTES")
        headers = [
            "RUC", "CLIENTE", "NOMBRE CONTACTO", "CORREO PRINCIPAL", "CORREO CC",
            "ENVIAR RECORDATORIOS", "BLOQUEADO", "DÍAS ANTICIPACIÓN", "FRECUENCIA DÍAS",
            "ÚLTIMO CORREO", "OBSERVACIONES"
        ]
        ws.append(headers)
        for c in range(1, len(headers) + 1):
            ws.cell(1, c).fill = PatternFill("solid", fgColor=BLUE)
            ws.cell(1, c).font = Font(color=WHITE, bold=True)
        ws.freeze_panes = "A2"
        widths = [16, 38, 24, 34, 34, 22, 14, 18, 18, 20, 45]
        for idx, width in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + idx)].width = width
    if "HISTORIAL_CORREOS" not in wb.sheetnames:
        ws = wb.create_sheet("HISTORIAL_CORREOS")
        headers = [
            "FECHA", "CLAVE_ENVÍO", "RUC", "CLIENTE", "DESTINATARIO", "CC",
            "COMPROBANTES", "TOTAL", "PRIORIDAD", "ESTADO", "ASUNTO", "ARCHIVO ORIGEN"
        ]
        ws.append(headers)
        for c in range(1, len(headers) + 1):
            ws.cell(1, c).fill = PatternFill("solid", fgColor=BLUE)
            ws.cell(1, c).font = Font(color=WHITE, bold=True)
        ws.freeze_panes = "A2"
        widths = [20, 55, 16, 38, 34, 34, 55, 16, 14, 20, 55, 45]
        from openpyxl.utils import get_column_letter
        for idx, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = width


def sync_master_clients(wb, records: List[Dict[str, Any]]) -> None:
    """Añade al maestro los clientes que aún no existen, sin sobrescribir correos."""
    ensure_email_sheets(wb)
    ws = wb["MAESTRO_CLIENTES"]
    cols = _headers(ws)
    existing_ruc = set()
    existing_name = set()
    for r in range(2, ws.max_row + 1):
        if "RUC" in cols:
            value = _ruc(ws.cell(r, cols["RUC"]).value)
            if value:
                existing_ruc.add(value)
        if "CLIENTE" in cols:
            value = _norm(ws.cell(r, cols["CLIENTE"]).value)
            if value:
                existing_name.add(value)

    unique: Dict[str, Dict[str, Any]] = {}
    for rec in records:
        ruc = _ruc(rec.get("RUC", ""))
        client = _text(rec.get("CLIENTE", ""))
        if not ruc and not client:
            continue
        key = ruc or _norm(client)
        unique.setdefault(key, {
            "ruc": ruc,
            "client": client,
            "responsable": _text(rec.get("RESPONSABLE", "")),
        })

    for item in sorted(unique.values(), key=lambda x: (x["client"].upper(), x["ruc"])):
        if (item["ruc"] and item["ruc"] in existing_ruc) or _norm(item["client"]) in existing_name:
            continue
        ws.append([
            item["ruc"], item["client"], "", "", "", "SI", "NO", 7, 3, "",
            f"Responsable actual: {item['responsable']}" if item["responsable"] else ""
        ])
        if item["ruc"]:
            existing_ruc.add(item["ruc"])
        existing_name.add(_norm(item["client"]))


def _headers(ws) -> Dict[str, int]:
    result: Dict[str, int] = {}
    for cell in ws[1]:
        result[_norm(cell.value)] = cell.column
    return result


def _load_master(ws) -> Tuple[Dict[str, Dict[str, Any]], Dict[str, Dict[str, Any]]]:
    cols = _headers(ws)
    by_ruc: Dict[str, Dict[str, Any]] = {}
    by_name: Dict[str, Dict[str, Any]] = {}
    for row in range(2, ws.max_row + 1):
        get = lambda name: ws.cell(row, cols[name]).value if name in cols else ""
        ruc = _ruc(get("RUC"))
        client = _text(get("CLIENTE"))
        if not ruc and not client:
            continue
        record = {
            "row": row,
            "ruc": ruc,
            "client": client,
            "contact": _text(get("NOMBRE CONTACTO")),
            "to": _split_emails(get("CORREO PRINCIPAL")),
            "cc": _split_emails(get("CORREO CC")),
            "enabled": _yes(get("ENVIAR RECORDATORIOS"), True),
            "blocked": _yes(get("BLOQUEADO"), False),
            "advance": int(get("DÍAS ANTICIPACIÓN") or 7),
            "frequency": int(get("FRECUENCIA DÍAS") or 3),
        }
        if ruc:
            by_ruc[ruc] = record
        if client:
            by_name[_norm(client)] = record
    return by_ruc, by_name


def _read_reminders(ws) -> List[Dict[str, Any]]:
    cols = _headers(ws)
    required = ["PRIORIDAD", "RUC", "CLIENTE", "COMPROBANTE", "SALDO", "VENCIMIENTO"]
    missing = [x for x in required if x not in cols]
    if missing:
        raise ValueError("Faltan columnas en RECORDATORIOS: " + ", ".join(missing))
    rows: List[Dict[str, Any]] = []
    for r in range(2, ws.max_row + 1):
        priority = _norm(ws.cell(r, cols["PRIORIDAD"]).value)
        if priority not in {"CRÍTICA", "CRITICA", "VENCIDA", "ALTA", "PREVENTIVA"}:
            continue
        saldo = ws.cell(r, cols["SALDO"]).value or 0
        try:
            saldo = float(saldo)
        except (TypeError, ValueError):
            saldo = 0.0
        rows.append({
            "priority": priority,
            "ruc": _ruc(ws.cell(r, cols["RUC"]).value),
            "client": _text(ws.cell(r, cols["CLIENTE"]).value),
            "document": _text(ws.cell(r, cols["COMPROBANTE"]).value),
            "amount": saldo,
            "due": ws.cell(r, cols["VENCIMIENTO"]).value,
            "days": ws.cell(r, cols.get("DÍAS", 0)).value if cols.get("DÍAS") else "",
        })
    return rows


def _priority(items: List[Dict[str, Any]]) -> str:
    rank = {"CRÍTICA": 0, "CRITICA": 0, "VENCIDA": 1, "ALTA": 2, "PREVENTIVA": 3}
    return min((x["priority"] for x in items), key=lambda p: rank.get(p, 99))


def _format_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return _text(value)


def _build_html(client: str, contact: str, items: List[Dict[str, Any]], total: float, signature: str) -> str:
    greeting = f"Estimado/a {html.escape(contact)}:" if contact else "Estimados señores:"
    rows = []
    for item in sorted(items, key=lambda x: (_format_date(x["due"]), x["document"])):
        rows.append(
            "<tr>"
            f"<td style='padding:6px 10px;border:1px solid #d9d9d9'>{html.escape(item['document'])}</td>"
            f"<td style='padding:6px 10px;border:1px solid #d9d9d9'>{html.escape(_format_date(item['due']))}</td>"
            f"<td style='padding:6px 10px;border:1px solid #d9d9d9;text-align:right'>S/ {item['amount']:,.2f}</td>"
            "</tr>"
        )
    table_rows = "".join(rows)
    return f"""
<html><body style="font-family:Calibri,Arial,sans-serif;font-size:11pt;color:#222">
<p>{greeting}</p>
<p>Les recordamos que {html.escape(client)} mantiene los siguientes comprobantes pendientes de cancelación:</p>
<table style="border-collapse:collapse;margin:12px 0">
<tr style="background:#1f4e78;color:white;font-weight:bold">
<td style="padding:6px 10px;border:1px solid #d9d9d9">Comprobante</td>
<td style="padding:6px 10px;border:1px solid #d9d9d9">Vencimiento</td>
<td style="padding:6px 10px;border:1px solid #d9d9d9">Saldo pendiente</td>
</tr>{table_rows}
<tr style="font-weight:bold;background:#ddebf7">
<td colspan="2" style="padding:6px 10px;border:1px solid #d9d9d9">Total pendiente</td>
<td style="padding:6px 10px;border:1px solid #d9d9d9;text-align:right">S/ {total:,.2f}</td>
</tr>
</table>
<p>Agradeceremos confirmar la fecha programada de pago o remitir la constancia correspondiente.</p>
<p>Saludos cordiales,<br>{html.escape(signature)}</p>
</body></html>
"""



def outlook_classic_available() -> Tuple[bool, str]:
    """Detecta Outlook clásico de escritorio. El nuevo Outlook no expone COM."""
    if os.name != "nt":
        return False, "Esta función solo está disponible en Windows."

    registry_paths = [
        (winreg.HKEY_CURRENT_USER, r"SOFTWARE\Microsoft\Windows\CurrentVersion\App Paths\OUTLOOK.EXE"),
        (winreg.HKEY_LOCAL_MACHINE, r"SOFTWARE\Microsoft\Windows\CurrentVersion\App Paths\OUTLOOK.EXE"),
        (winreg.HKEY_LOCAL_MACHINE, r"SOFTWARE\WOW6432Node\Microsoft\Windows\CurrentVersion\App Paths\OUTLOOK.EXE"),
    ]
    for hive, key_path in registry_paths:
        try:
            with winreg.OpenKey(hive, key_path) as key:
                exe, _ = winreg.QueryValueEx(key, None)
                if exe and Path(exe).exists():
                    return True, str(exe)
        except OSError:
            pass

    return False, (
        "No se encontró Outlook clásico de escritorio. "
        "El nuevo Outlook, Outlook web y Gmail no permiten crear borradores mediante COM."
    )


def _prepare_email_groups(workbook_path: Path):
    path = Path(workbook_path)
    if not path.is_file():
        raise FileNotFoundError(f"No existe el archivo: {path}")

    wb = load_workbook(path)
    ensure_email_sheets(wb)
    if "RECORDATORIOS" not in wb.sheetnames:
        wb.close()
        raise ValueError("El archivo no contiene la hoja RECORDATORIOS. Actualice los créditos primero.")

    master_ws = wb["MAESTRO_CLIENTES"]
    by_ruc, by_name = _load_master(master_ws)
    reminders = _read_reminders(wb["RECORDATORIOS"])
    groups: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for item in reminders:
        key = item["ruc"] or _norm(item["client"])
        groups[key].append(item)
    return wb, master_ws, by_ruc, by_name, groups


def generate_html_previews(
    workbook_path: Path,
    output_dir: Path | None = None,
    signature: str = "Área de Cobranzas",
) -> Dict[str, Any]:
    """Genera vistas previas HTML sin requerir Outlook."""
    path = Path(workbook_path)
    output_dir = Path(output_dir) if output_dir else path.parent / f"VISTA_PREVIA_CORREOS_{datetime.now():%Y%m%d_%H%M%S}"
    output_dir.mkdir(parents=True, exist_ok=True)

    wb, _master_ws, by_ruc, by_name, groups = _prepare_email_groups(path)
    created = 0
    missing_email: List[str] = []
    blocked: List[str] = []
    errors: List[str] = []
    summary_rows = []

    for _group_key, items in groups.items():
        client = items[0]["client"]
        ruc = items[0]["ruc"]
        master = by_ruc.get(ruc) if ruc else None
        if not master:
            master = by_name.get(_norm(client))
        if not master or not master["to"]:
            missing_email.append(f"{ruc} - {client}".strip(" -"))
            continue
        if not master["enabled"] or master["blocked"]:
            blocked.append(f"{ruc} - {client}".strip(" -"))
            continue

        documents = sorted(x["document"] for x in items)
        total = sum(x["amount"] for x in items)
        priority = _priority(items)
        subject = f"Recordatorio de comprobantes pendientes - {client}"
        safe_name = re.sub(r'[^A-Za-z0-9ÁÉÍÓÚÜÑáéíóúüñ _.-]+', '_', client).strip()[:80] or "cliente"
        file_path = output_dir / f"{created + 1:03d}_{safe_name}.html"
        try:
            body = _build_html(client, master["contact"], items, total, signature)
            metadata = (
                "<!doctype html><html><head><meta charset='utf-8'>"
                f"<title>{html.escape(subject)}</title></head><body>"
                f"<p><strong>Para:</strong> {html.escape('; '.join(master['to']))}<br>"
                f"<strong>CC:</strong> {html.escape('; '.join(master['cc']))}<br>"
                f"<strong>Asunto:</strong> {html.escape(subject)}</p><hr>"
            )
            # body already contains html/body tags, extract inner content for valid document
            inner = re.sub(r'^\s*<html><body[^>]*>|</body></html>\s*$', '', body.strip(), flags=re.I)
            file_path.write_text(metadata + inner + "</body></html>", encoding="utf-8")
            created += 1
            summary_rows.append([
                ruc, client, "; ".join(master["to"]), "; ".join(master["cc"]),
                ", ".join(documents), f"{total:.2f}", priority, subject, file_path.name,
            ])
        except Exception as exc:
            errors.append(f"{client}: {exc}")

    with (output_dir / "RESUMEN_CORREOS.csv").open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh, delimiter=';')
        writer.writerow(["RUC", "CLIENTE", "PARA", "CC", "COMPROBANTES", "TOTAL", "PRIORIDAD", "ASUNTO", "ARCHIVO"])
        writer.writerows(summary_rows)

    wb.close()
    return {
        "created": created,
        "missing_email": missing_email,
        "blocked": blocked,
        "errors": errors,
        "output_dir": str(output_dir),
    }


def generate_outlook_drafts(workbook_path: Path, signature: str = "Área de Cobranzas") -> Dict[str, Any]:
    available, detail = outlook_classic_available()
    if not available:
        raise RuntimeError(
            detail + "\n\nUse 'Generar vista previa HTML' en esta computadora, "
            "o ejecute los borradores en una PC con Outlook clásico configurado."
        )

    try:
        import pythoncom
        import win32com.client as win32
    except ImportError as exc:
        raise RuntimeError(
            "Falta el componente de Outlook. Vuelva a generar el EXE con pywin32 incluido."
        ) from exc

    pythoncom.CoInitialize()
    path = Path(workbook_path)
    if not path.is_file():
        raise FileNotFoundError(f"No existe el archivo: {path}")

    wb = load_workbook(path)
    ensure_email_sheets(wb)
    if "RECORDATORIOS" not in wb.sheetnames:
        wb.close()
        raise ValueError("El archivo no contiene la hoja RECORDATORIOS. Actualice los créditos primero.")

    master_ws = wb["MAESTRO_CLIENTES"]
    hist_ws = wb["HISTORIAL_CORREOS"]
    by_ruc, by_name = _load_master(master_ws)
    reminders = _read_reminders(wb["RECORDATORIOS"])

    previous_keys = set()
    hist_cols = _headers(hist_ws)
    key_col = hist_cols.get("CLAVE ENVÍO") or hist_cols.get("CLAVE ENVIO")
    if key_col:
        for r in range(2, hist_ws.max_row + 1):
            previous_keys.add(_text(hist_ws.cell(r, key_col).value))

    groups: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for item in reminders:
        key = item["ruc"] or _norm(item["client"])
        groups[key].append(item)

    try:
        outlook = win32.GetActiveObject("Outlook.Application")
    except Exception:
        outlook = win32.Dispatch("Outlook.Application")
    try:
        namespace = outlook.GetNamespace("MAPI")
        if namespace.Folders.Count == 0:
            raise RuntimeError("Outlook no tiene un perfil de correo configurado.")
    except Exception as exc:
        wb.close()
        pythoncom.CoUninitialize()
        raise RuntimeError(
            "Outlook clásico está instalado, pero no hay un perfil de correo disponible. "
            "Abra Outlook, configure/inicie sesión y vuelva a intentar."
        ) from exc

    created = 0
    duplicates = 0
    missing_email: List[str] = []
    blocked: List[str] = []
    errors: List[str] = []
    today = date.today().isoformat()

    for group_key, items in groups.items():
        client = items[0]["client"]
        ruc = items[0]["ruc"]
        master = by_ruc.get(ruc) if ruc else None
        if not master:
            master = by_name.get(_norm(client))
        if not master or not master["to"]:
            missing_email.append(f"{ruc} - {client}".strip(" -"))
            continue
        if not master["enabled"] or master["blocked"]:
            blocked.append(f"{ruc} - {client}".strip(" -"))
            continue

        documents = sorted(x["document"] for x in items)
        send_key = f"{today}|{ruc or _norm(client)}|{'|'.join(documents)}"
        if send_key in previous_keys:
            duplicates += 1
            continue

        total = sum(x["amount"] for x in items)
        priority = _priority(items)
        subject = f"Recordatorio de comprobantes pendientes - {client}"
        try:
            mail = outlook.CreateItem(0)
            mail.To = "; ".join(master["to"])
            mail.CC = "; ".join(master["cc"])
            mail.Subject = subject
            mail.HTMLBody = _build_html(client, master["contact"], items, total, signature)
            mail.Save()
            hist_ws.append([
                datetime.now(), send_key, ruc, client, mail.To, mail.CC,
                ", ".join(documents), total, priority, "BORRADOR CREADO", subject, path.name
            ])
            created += 1
            previous_keys.add(send_key)
            if "ÚLTIMO CORREO" in _headers(master_ws):
                master_ws.cell(master["row"], _headers(master_ws)["ÚLTIMO CORREO"]).value = datetime.now()
        except Exception as exc:
            errors.append(f"{client}: {exc}")

    for r in range(2, hist_ws.max_row + 1):
        hist_ws.cell(r, 1).number_format = "dd/mm/yyyy hh:mm"
        hist_ws.cell(r, 8).number_format = "#,##0.00"
    wb.save(path)
    wb.close()
    pythoncom.CoUninitialize()

    return {
        "created": created,
        "duplicates": duplicates,
        "missing_email": missing_email,
        "blocked": blocked,
        "errors": errors,
        "workbook": str(path),
    }
