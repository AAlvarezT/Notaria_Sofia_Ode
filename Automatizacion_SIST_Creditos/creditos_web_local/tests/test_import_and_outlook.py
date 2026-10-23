from __future__ import annotations

import sqlite3
import tempfile
from pathlib import Path

import pytest
from openpyxl import Workbook

from app import is_valid_email
from database import (
    connect,
    draft_already_exists,
    init_db,
    list_documents,
    mark_draft_created,
    mark_draft_failed,
    reserve_draft_slot,
    save_draft_history,
)
from outlook import MockOutlookDraftService, OutlookDraftService
from scheduler import _check_folder
from services import determine_document_state, import_signo_file


@pytest.fixture
def temp_db(tmp_path):
    db_path = tmp_path / "creditos.sqlite3"
    init_db(db_path)
    return db_path


def _write_signo_file(path: Path, rows_by_sheet: dict, include_title_row: bool = False):
    """Create a minimal SIGNO-style Excel file.

    rows_by_sheet: {sheet_name: [header_row, data_row, ...]}
    include_title_row: if True, inserts a dummy title row before the headers.
    """
    wb = Workbook()
    # Remove default sheet
    for s in list(wb.sheetnames):
        wb.remove(wb[s])

    for sheet_name, rows in rows_by_sheet.items():
        ws = wb.create_sheet(title=sheet_name)
        if include_title_row:
            ws.append([f"Reporte {sheet_name} - 2024-10-21"])  # title row
        for row in rows:
            ws.append(row)

    wb.save(path)


# ---------------------------------------------------------------------------
# Existing tests (preserved and updated for new behavior)
# ---------------------------------------------------------------------------

def test_import_new_format_keeps_leading_zeroes(temp_db, tmp_path):
    path = tmp_path / "signo.xlsx"
    _write_signo_file(
        path,
        {
            "Pendientes": [
                ["#", "DOI", "Cliente", "FechaEmision", "Tipo", "Serie", "Numero", "Kardex", "ImporteFacturado", "SaldoPendiente", "FechaVencimiento", "Empleado", "FechaActualizacion"],
                [1, "12345678", "Cliente Demo", "2024-01-01", "F", "001", "0001", "K", "102.50", "52.50", "2024-02-05", "Ana", "2024-02-02"],
            ]
        },
    )
    result = import_signo_file(path, temp_db)
    assert result["status"] == "OK"
    with connect(temp_db) as conn:
        row = conn.execute("SELECT serie, numero, saldo_observado, estado FROM documents WHERE document_key = ?", ("F|001|0001",)).fetchone()
    assert row["serie"] == "001"
    assert row["numero"] == "0001"
    assert row["saldo_observado"] == "52.50"


def test_partial_balance_and_document_present_in_pending_and_cancelled(temp_db, tmp_path):
    path = tmp_path / "signo.xlsx"
    _write_signo_file(
        path,
        {
            "Pendientes": [
                ["#", "DOI", "Cliente", "FechaEmision", "Tipo", "Serie", "Numero", "Kardex", "ImporteFacturado", "SaldoPendiente", "FechaVencimiento", "Empleado", "FechaActualizacion"],
                [1, "201", "Cliente A", "2024-01-01", "F", "B1", "10", "K", "500", "200", "2024-02-05", "Ana", "2024-02-02"],
            ],
            "Cancelados": [
                ["#", "DOI", "Cliente", "FechaEmision", "Tipo", "Serie", "Numero", "Kardex", "ImporteFacturado", "SaldoPendiente", "FormaPago", "FechaCancelacion", "FechaVencimiento", "Empleado", "FechaActualizacion"],
                [1, "201", "Cliente A", "2024-01-01", "F", "B1", "10", "K", "500", "200", "EFECTIVO", "2024-01-20", "2024-02-05", "Ana", "2024-02-02"],
            ],
        },
    )
    import_signo_file(path, temp_db)
    with connect(temp_db) as conn:
        row = conn.execute("SELECT estado, saldo_observado FROM documents WHERE document_key = ?", ("F|B1|10",)).fetchone()
    assert row["estado"] == "SALDO_PARCIAL"
    assert row["saldo_observado"] == "200"


def test_zero_balance_document_is_excluded(temp_db, tmp_path):
    path = tmp_path / "signo.xlsx"
    _write_signo_file(
        path,
        {
            "Pendientes": [
                ["#", "DOI", "Cliente", "FechaEmision", "Tipo", "Serie", "Numero", "Kardex", "ImporteFacturado", "SaldoPendiente", "FechaVencimiento", "Empleado", "FechaActualizacion"],
                [1, "202", "Cliente B", "2024-01-01", "F", "B2", "11", "K", "150", "0", "2024-02-05", "Ana", "2024-02-02"],
            ]
        },
    )
    import_signo_file(path, temp_db)
    with connect(temp_db) as conn:
        row = conn.execute("SELECT estado FROM documents WHERE document_key = ?", ("F|B2|11",)).fetchone()
    assert row["estado"] == "EXCLUIDO"


def test_missing_document_is_not_marked_as_paid(temp_db, tmp_path):
    path = tmp_path / "signo.xlsx"
    _write_signo_file(
        path,
        {
            "Ventas": [
                ["#", "DOI", "Cliente", "FechaEmision", "Tipo", "Serie", "Numero", "Kardex", "ImporteFacturado", "SaldoPendiente", "FormaPago", "FechaVencimiento", "Empleado", "FechaActualizacion"],
                [1, "303", "Cliente C", "2024-01-01", "F", "B3", "12", "K", "200", "200", "EFECTIVO", "2024-02-05", "Ana", "2024-02-02"],
            ]
        },
    )
    import_signo_file(path, temp_db)
    with connect(temp_db) as conn:
        row = conn.execute("SELECT estado FROM documents WHERE document_key = ?", ("F|B3|12",)).fetchone()
    assert row["estado"] == "EN_REVISION"


def test_repeated_import_is_idempotent_and_older_report_does_not_overwrite(temp_db, tmp_path):
    path = tmp_path / "signo.xlsx"
    _write_signo_file(
        path,
        {
            "Pendientes": [
                ["#", "DOI", "Cliente", "FechaEmision", "Tipo", "Serie", "Numero", "Kardex", "ImporteFacturado", "SaldoPendiente", "FechaVencimiento", "Empleado", "FechaActualizacion"],
                [1, "444", "Cliente D", "2024-01-01", "F", "B4", "13", "K", "100", "100", "2024-02-05", "Ana", "2024-02-01"],
            ]
        },
    )
    first = import_signo_file(path, temp_db)
    second = import_signo_file(path, temp_db)
    assert first["documents"] == 1
    assert second["documents"] == 1
    with connect(temp_db) as conn:
        count = conn.execute("SELECT COUNT(*) FROM documents WHERE document_key = ?", ("F|B4|13",)).fetchone()[0]
    assert count == 1


def test_older_report_does_not_overwrite_newer(temp_db, tmp_path):
    path_old = tmp_path / "old.xlsx"
    _write_signo_file(
        path_old,
        {
            "Pendientes": [
                ["#", "DOI", "Cliente", "FechaEmision", "Tipo", "Serie", "Numero", "Kardex", "ImporteFacturado", "SaldoPendiente", "FechaVencimiento", "Empleado", "FechaActualizacion"],
                [1, "555", "Cliente E", "2024-01-01", "F", "B5", "14", "K", "100", "100", "2024-02-05", "Ana", "2024-01-01"],
            ]
        },
    )
    path_new = tmp_path / "new.xlsx"
    _write_signo_file(
        path_new,
        {
            "Pendientes": [
                ["#", "DOI", "Cliente", "FechaEmision", "Tipo", "Serie", "Numero", "Kardex", "ImporteFacturado", "SaldoPendiente", "FechaVencimiento", "Empleado", "FechaActualizacion"],
                [1, "555", "Cliente E", "2024-01-01", "F", "B5", "14", "K", "100", "60", "2024-02-05", "Ana", "2024-02-02"],
            ]
        },
    )
    import_signo_file(path_old, temp_db)
    import_signo_file(path_new, temp_db)
    with connect(temp_db) as conn:
        row = conn.execute("SELECT saldo_observado, fecha_actualizacion FROM documents WHERE document_key = ?", ("F|B5|14",)).fetchone()
    assert row["saldo_observado"] == "60"
    assert row["fecha_actualizacion"] == "2024-02-02T00:00:00"


def test_manual_information_is_preserved_when_same_document_is_updated(temp_db, tmp_path):
    path = tmp_path / "signo.xlsx"
    _write_signo_file(
        path,
        {
            "Pendientes": [
                ["#", "DOI", "Cliente", "FechaEmision", "Tipo", "Serie", "Numero", "Kardex", "ImporteFacturado", "SaldoPendiente", "FechaVencimiento", "Empleado", "FechaActualizacion"],
                [1, "666", "Cliente F", "2024-01-01", "F", "B6", "15", "K", "100", "100", "2024-02-05", "Ana", "2024-02-02"],
            ]
        },
    )
    import_signo_file(path, temp_db)
    with connect(temp_db) as conn:
        conn.execute(
            "UPDATE documents SET observaciones = 'Compromiso manual', fecha_vencimiento = '2024-03-07' WHERE document_key = ?",
            ("F|B6|15",),
        )
    import_signo_file(path, temp_db)
    with connect(temp_db) as conn:
        row = conn.execute("SELECT observaciones, fecha_vencimiento FROM documents WHERE document_key = ?", ("F|B6|15",)).fetchone()
    assert row["observaciones"] == "Compromiso manual"
    assert row["fecha_vencimiento"] == "2024-03-07"


def test_frequency_anticipation_and_deduplication_work_for_outlook_mocks():
    service = MockOutlookDraftService()
    result_a = service.create_draft("Recordatorio", "<p>Body</p>", "cliente@demo.com", dedupe_key="cli-1")
    result_b = service.create_draft("Recordatorio", "<p>Body</p>", "cliente@demo.com", dedupe_key="cli-1")
    assert result_a.success is True
    assert result_b.success is True
    assert "DUPLICATE" in result_b.message_id
    service.release()
    assert service.released is True


def test_outlook_failure_is_not_accepted_as_real_success():
    service = OutlookDraftService(enabled=True)
    result = service.create_draft("Recordatorio", "<p>Body</p>", "cliente@demo.com", dedupe_key="x")
    assert result.success is False or "DUPLICATE" in (result.message_id or "") or result.error is not None


def test_document_state_rules():
    assert determine_document_state("F|1|1", {"saldo_pendiente": "0", "importe_facturado": "100", "source_sheet": "Pendientes"}) == "EXCLUIDO"
    assert determine_document_state("F|1|2", {"saldo_pendiente": "150", "importe_facturado": "200", "source_sheet": "Pendientes"}) == "SALDO_PARCIAL"
    assert determine_document_state("F|1|3", {"saldo_pendiente": "50", "importe_facturado": "0", "source_sheet": "Ventas"}) == "EN_REVISION"


# ---------------------------------------------------------------------------
# New tests
# ---------------------------------------------------------------------------

def test_title_row_plus_header_row(temp_db, tmp_path):
    """Files with a title row (row 1) and headers (row 2) are imported correctly."""
    path = tmp_path / "signo_title.xlsx"
    _write_signo_file(
        path,
        {
            "Pendientes": [
                ["#", "DOI", "Cliente", "FechaEmision", "Tipo", "Serie", "Numero", "Kardex", "ImporteFacturado", "SaldoPendiente", "FechaVencimiento", "Empleado", "FechaActualizacion"],
                [1, "99999999", "Empresa XYZ", "2024-03-01", "F", "FE01", "0030712", "K", "5000", "1320", "2024-04-01", "Luis", "2024-10-21 13:34:48"],
            ]
        },
        include_title_row=True,
    )
    result = import_signo_file(path, temp_db)
    assert result["status"] == "OK"
    with connect(temp_db) as conn:
        row = conn.execute("SELECT * FROM documents WHERE document_key = ?", ("F|FE01|0030712",)).fetchone()
    assert row is not None, "Documento no importado"
    assert row["serie"] == "FE01"
    assert row["numero"] == "0030712"
    assert row["saldo_observado"] == "1320"


def test_partial_balance_in_all_three_sheets(temp_db, tmp_path):
    """F|FE01|0030712 appearing in Ventas+Pendientes+Cancelados keeps SALDO_PARCIAL."""
    path = tmp_path / "tres_hojas.xlsx"
    _write_signo_file(
        path,
        {
            "Ventas": [
                ["#", "DOI", "Cliente", "FechaEmision", "Tipo", "Serie", "Numero", "Kardex", "ImporteFacturado", "SaldoPendiente", "FormaPago", "FechaVencimiento", "Empleado", "FechaActualizacion"],
                [1, "99999999", "Empresa XYZ", "2024-03-01", "F", "FE01", "0030712", "K", "5000", "1320", "CREDITO", "2024-04-01", "Luis", "2024-10-21 13:34:48"],
            ],
            "Pendientes": [
                ["#", "DOI", "Cliente", "FechaEmision", "Tipo", "Serie", "Numero", "Kardex", "ImporteFacturado", "SaldoPendiente", "FechaVencimiento", "Empleado", "FechaActualizacion"],
                [1, "99999999", "Empresa XYZ", "2024-03-01", "F", "FE01", "0030712", "K", "5000", "1320", "2024-04-01", "Luis", "2024-10-21 13:34:48"],
            ],
            "Cancelados": [
                ["#", "DOI", "Cliente", "FechaEmision", "Tipo", "Serie", "Numero", "Kardex", "ImporteFacturado", "SaldoPendiente", "FormaPago", "FechaCancelacion", "FechaVencimiento", "Empleado", "FechaActualizacion"],
                [1, "99999999", "Empresa XYZ", "2024-03-01", "F", "FE01", "0030712", "K", "5000", "3680", "DETRACCION", "2024-09-15", "2024-04-01", "Luis", "2024-10-21 13:34:48"],
            ],
        },
    )
    import_signo_file(path, temp_db)
    with connect(temp_db) as conn:
        row = conn.execute("SELECT estado, saldo_observado FROM documents WHERE document_key = ?", ("F|FE01|0030712",)).fetchone()
    assert row is not None
    # Pendientes is authoritative source → saldo = 1320
    assert row["saldo_observado"] == "1320"
    # saldo 1320 < importe 5000 → SALDO_PARCIAL
    assert row["estado"] == "SALDO_PARCIAL"


def test_saldo_parcial_has_no_revision_required(temp_db, tmp_path):
    """SALDO_PARCIAL alone must not set revision_required=1 (it is a valid state)."""
    path = tmp_path / "parcial.xlsx"
    _write_signo_file(
        path,
        {
            "Pendientes": [
                ["#", "DOI", "Cliente", "FechaEmision", "Tipo", "Serie", "Numero", "Kardex", "ImporteFacturado", "SaldoPendiente", "FechaVencimiento", "Empleado", "FechaActualizacion"],
                [1, "11111111", "Emp Parcial", "2024-01-01", "F", "C1", "01", "K", "1000", "400", "2024-06-01", "Ana", "2024-06-01"],
            ]
        },
    )
    import_signo_file(path, temp_db)
    with connect(temp_db) as conn:
        row = conn.execute("SELECT estado, revision_required FROM documents WHERE document_key = ?", ("F|C1|01",)).fetchone()
    assert row["estado"] == "SALDO_PARCIAL"
    assert row["revision_required"] == 0


def test_partial_balance_included_in_cobranza_list(temp_db, tmp_path):
    """list_documents({'estado': ['PENDIENTE', 'SALDO_PARCIAL']}) includes SALDO_PARCIAL."""
    path = tmp_path / "cobrar.xlsx"
    _write_signo_file(
        path,
        {
            "Pendientes": [
                ["#", "DOI", "Cliente", "FechaEmision", "Tipo", "Serie", "Numero", "Kardex", "ImporteFacturado", "SaldoPendiente", "FechaVencimiento", "Empleado", "FechaActualizacion"],
                [1, "22222222", "Emp A", "2024-01-01", "F", "D1", "01", "K", "500", "500", "2024-06-01", "Ana", "2024-06-01"],
                [2, "33333333", "Emp B", "2024-01-01", "F", "D2", "02", "K", "800", "300", "2024-07-01", "Ana", "2024-06-01"],
            ]
        },
    )
    import_signo_file(path, temp_db)
    docs = list_documents({"estado": ["PENDIENTE", "SALDO_PARCIAL"]}, db_path=temp_db)
    keys = {d["document_key"] for d in docs}
    assert "F|D1|01" in keys, "PENDIENTE no está en cobranza"
    assert "F|D2|02" in keys, "SALDO_PARCIAL no está en cobranza"


def test_date_normalized_to_iso_and_timestamp_preserves_time(temp_db, tmp_path):
    """fecha_emision → ISO date; fecha_actualizacion → ISO datetime with time."""
    path = tmp_path / "fechas.xlsx"
    _write_signo_file(
        path,
        {
            "Pendientes": [
                ["#", "DOI", "Cliente", "FechaEmision", "Tipo", "Serie", "Numero", "Kardex", "ImporteFacturado", "SaldoPendiente", "FechaVencimiento", "Empleado", "FechaActualizacion"],
                [1, "77777777", "Emp F", "01/03/2024", "F", "E1", "01", "K", "200", "200", "05/04/2024", "Ana", "21/10/2024 13:34:48"],
            ]
        },
    )
    import_signo_file(path, temp_db)
    with connect(temp_db) as conn:
        row = conn.execute("SELECT fecha_emision, fecha_vencimiento, fecha_actualizacion FROM documents WHERE document_key = ?", ("F|E1|01",)).fetchone()
    assert row["fecha_emision"] == "2024-03-01"
    assert row["fecha_vencimiento"] == "2024-04-05"
    assert row["fecha_actualizacion"] == "2024-10-21T13:34:48"


def test_same_day_update_different_hours_newer_wins(temp_db, tmp_path):
    """When two imports share the same date but different times, the later time wins."""
    path_early = tmp_path / "early.xlsx"
    _write_signo_file(
        path_early,
        {
            "Pendientes": [
                ["#", "DOI", "Cliente", "FechaEmision", "Tipo", "Serie", "Numero", "Kardex", "ImporteFacturado", "SaldoPendiente", "FechaVencimiento", "Empleado", "FechaActualizacion"],
                [1, "44444444", "Emp G", "2024-01-01", "F", "G1", "01", "K", "300", "300", "2024-06-01", "Ana", "21/10/2024 08:00:00"],
            ]
        },
    )
    path_late = tmp_path / "late.xlsx"
    _write_signo_file(
        path_late,
        {
            "Pendientes": [
                ["#", "DOI", "Cliente", "FechaEmision", "Tipo", "Serie", "Numero", "Kardex", "ImporteFacturado", "SaldoPendiente", "FechaVencimiento", "Empleado", "FechaActualizacion"],
                [1, "44444444", "Emp G", "2024-01-01", "F", "G1", "01", "K", "300", "150", "2024-06-01", "Ana", "21/10/2024 15:30:00"],
            ]
        },
    )
    import_signo_file(path_early, temp_db)
    import_signo_file(path_late, temp_db)
    with connect(temp_db) as conn:
        row = conn.execute("SELECT saldo_observado, fecha_actualizacion FROM documents WHERE document_key = ?", ("F|G1|01",)).fetchone()
    assert row["saldo_observado"] == "150", "El reporte más reciente del mismo día debe ganar"
    assert "15:30:00" in row["fecha_actualizacion"]


def test_reimport_no_duplicates(temp_db, tmp_path):
    """Re-importing the same file must not create duplicate document rows."""
    path = tmp_path / "reimport.xlsx"
    _write_signo_file(
        path,
        {
            "Pendientes": [
                ["#", "DOI", "Cliente", "FechaEmision", "Tipo", "Serie", "Numero", "Kardex", "ImporteFacturado", "SaldoPendiente", "FechaVencimiento", "Empleado", "FechaActualizacion"],
                [1, "55555555", "Emp H", "2024-01-01", "F", "H1", "01", "K", "100", "100", "2024-06-01", "Ana", "2024-06-01"],
                [2, "66666666", "Emp I", "2024-01-01", "F", "H2", "02", "K", "200", "200", "2024-07-01", "Ana", "2024-06-01"],
            ]
        },
    )
    import_signo_file(path, temp_db)
    import_signo_file(path, temp_db)
    with connect(temp_db) as conn:
        count = conn.execute("SELECT COUNT(*) FROM documents WHERE document_key LIKE 'F|H%'").fetchone()[0]
    assert count == 2


def test_old_report_does_not_overwrite_newer_data(temp_db, tmp_path):
    """Importing an older file after a newer one must not downgrade the saldo."""
    path_new = tmp_path / "new.xlsx"
    _write_signo_file(
        path_new,
        {
            "Pendientes": [
                ["#", "DOI", "Cliente", "FechaEmision", "Tipo", "Serie", "Numero", "Kardex", "ImporteFacturado", "SaldoPendiente", "FechaVencimiento", "Empleado", "FechaActualizacion"],
                [1, "77788800", "Emp J", "2024-01-01", "F", "J1", "01", "K", "500", "200", "2024-06-01", "Ana", "2024-10-20"],
            ]
        },
    )
    path_old = tmp_path / "old.xlsx"
    _write_signo_file(
        path_old,
        {
            "Pendientes": [
                ["#", "DOI", "Cliente", "FechaEmision", "Tipo", "Serie", "Numero", "Kardex", "ImporteFacturado", "SaldoPendiente", "FechaVencimiento", "Empleado", "FechaActualizacion"],
                [1, "77788800", "Emp J", "2024-01-01", "F", "J1", "01", "K", "500", "500", "2024-06-01", "Ana", "2024-09-01"],
            ]
        },
    )
    import_signo_file(path_new, temp_db)
    import_signo_file(path_old, temp_db)
    with connect(temp_db) as conn:
        row = conn.execute("SELECT saldo_observado FROM documents WHERE document_key = ?", ("F|J1|01",)).fetchone()
    assert row["saldo_observado"] == "200", "El reporte más antiguo no debe sobrescribir el más reciente"


def test_manual_contacts_preserved_on_reimport(temp_db, tmp_path):
    """Contacts added manually must survive a reimport of the same file."""
    path = tmp_path / "contact.xlsx"
    _write_signo_file(
        path,
        {
            "Pendientes": [
                ["#", "DOI", "Cliente", "FechaEmision", "Tipo", "Serie", "Numero", "Kardex", "ImporteFacturado", "SaldoPendiente", "FechaVencimiento", "Empleado", "FechaActualizacion"],
                [1, "88888888", "Emp K", "2024-01-01", "F", "K1", "01", "K", "100", "100", "2024-06-01", "Ana", "2024-06-01"],
            ]
        },
    )
    import_signo_file(path, temp_db)
    with connect(temp_db) as conn:
        conn.execute(
            "INSERT INTO contacts(document_key, nombre, correo, telefono, rol, es_principal)"
            " VALUES (?, ?, ?, ?, ?, ?)",
            ("F|K1|01", "Gerente", "gerente@emp.com", "999000111", "Pagos", 1),
        )
    import_signo_file(path, temp_db)
    with connect(temp_db) as conn:
        row = conn.execute(
            "SELECT correo FROM contacts WHERE document_key = ?", ("F|K1|01",)
        ).fetchone()
    assert row is not None and row["correo"] == "gerente@emp.com", "Contacto manual perdido tras reimportación"


def test_invalid_amount_sends_to_revision(temp_db, tmp_path):
    """A row with an invalid amount must be classified EN_REVISION, not EXCLUIDO."""
    path = tmp_path / "invalid_amt.xlsx"
    _write_signo_file(
        path,
        {
            "Pendientes": [
                ["#", "DOI", "Cliente", "FechaEmision", "Tipo", "Serie", "Numero", "Kardex", "ImporteFacturado", "SaldoPendiente", "FechaVencimiento", "Empleado", "FechaActualizacion"],
                [1, "99900099", "Emp L", "2024-01-01", "F", "L1", "01", "K", "N/D", "N/D", "2024-06-01", "Ana", "2024-06-01"],
            ]
        },
    )
    result = import_signo_file(path, temp_db)
    assert result["errors"], "Se esperaban errores de monto inválido"
    with connect(temp_db) as conn:
        row = conn.execute("SELECT estado FROM documents WHERE document_key = ?", ("F|L1|01",)).fetchone()
    assert row["estado"] == "EN_REVISION", "Monto inválido debe ir a revisión, no EXCLUIDO"


def test_outlook_disabled_returns_no_success():
    """OutlookDraftService with enabled=False must not return success=True."""
    service = OutlookDraftService(enabled=False)
    result = service.create_draft("Asunto", "<p>Cuerpo</p>", "cli@ejemplo.com")
    assert result.success is False
    assert result.error is not None
    assert "desactivad" in result.error.lower() or "deshabilitad" in result.error.lower()


def test_mock_outlook_dedup_in_session():
    """Mock dedup works within a session (in-memory _created_keys)."""
    svc = MockOutlookDraftService()
    r1 = svc.create_draft("Sub", "<p>B</p>", "a@b.com", dedupe_key="dk-1")
    r2 = svc.create_draft("Sub", "<p>B</p>", "a@b.com", dedupe_key="dk-1")
    assert r1.success is True
    assert "DUPLICATE" in r2.message_id


def test_persistent_draft_deduplication(temp_db):
    """draft_already_exists + save_draft_history provide cross-restart deduplication."""
    init_db(temp_db)  # already done by fixture; idempotent
    # Simulate writing the draft_history directly (as the /draft endpoint would)
    # First call – no duplicate
    assert not draft_already_exists("F|FE01|0001", "Recordatorio de pago", db_path=temp_db)
    save_draft_history(
        document_key="F|FE01|0001",
        cliente="Empresa XYZ",
        destinatario="pagos@empresa.com",
        subject="Recordatorio de pago",
        outlook_message_id="MOCK_DRAFT_123",
        draft_status="CREATED",
        html_preview="<p>body</p>",
        db_path=temp_db,
    )
    # Second call – now it IS a duplicate
    assert draft_already_exists("F|FE01|0001", "Recordatorio de pago", db_path=temp_db)


def test_cancelled_sheet_alone_does_not_pay_document(temp_db, tmp_path):
    """A document only in Cancelados with saldo>0 must NOT be EXCLUIDO."""
    path = tmp_path / "solo_cancelados.xlsx"
    _write_signo_file(
        path,
        {
            "Cancelados": [
                ["#", "DOI", "Cliente", "FechaEmision", "Tipo", "Serie", "Numero", "Kardex", "ImporteFacturado", "SaldoPendiente", "FormaPago", "FechaCancelacion", "FechaVencimiento", "Empleado", "FechaActualizacion"],
                [1, "12312312", "Emp M", "2024-01-01", "F", "M1", "01", "K", "1000", "600", "DETRACCION", "2024-06-01", "2024-06-30", "Ana", "2024-10-21"],
            ]
        },
    )
    import_signo_file(path, temp_db)
    with connect(temp_db) as conn:
        row = conn.execute("SELECT estado FROM documents WHERE document_key = ?", ("F|M1|01",)).fetchone()
    # saldo=600 < importe=1000 → SALDO_PARCIAL (not EXCLUIDO)
    assert row["estado"] in ("SALDO_PARCIAL", "EN_REVISION")
    assert row["estado"] != "EXCLUIDO"


def test_absent_date_0000_treated_as_null(temp_db, tmp_path):
    """The sentinel date 0000-00-00 must be stored as NULL, not as a string."""
    path = tmp_path / "null_date.xlsx"
    _write_signo_file(
        path,
        {
            "Pendientes": [
                ["#", "DOI", "Cliente", "FechaEmision", "Tipo", "Serie", "Numero", "Kardex", "ImporteFacturado", "SaldoPendiente", "FechaVencimiento", "Empleado", "FechaActualizacion"],
                [1, "10101010", "Emp N", "0000-00-00", "F", "N1", "01", "K", "200", "200", "0000-00-00", "Ana", "2024-06-01"],
            ]
        },
    )
    import_signo_file(path, temp_db)
    with connect(temp_db) as conn:
        row = conn.execute("SELECT fecha_emision, fecha_vencimiento FROM documents WHERE document_key = ?", ("F|N1|01",)).fetchone()
    assert row["fecha_emision"] is None
    assert row["fecha_vencimiento"] is None


def test_email_validation_is_stricter_than_at_only():
    assert is_valid_email("pagos@empresa.com")
    assert not is_valid_email("sin-arroba")
    assert not is_valid_email("nombre@localhost")
    assert not is_valid_email("invalido@@empresa.com")
    assert not is_valid_email("nombre @empresa.com")


def test_reserve_draft_slot_prevents_duplicates_and_allows_retry_after_failure(temp_db):
    reserved, status = reserve_draft_slot("F|A|1", "Cliente A", "a@demo.com", "Recordatorio", "<p>x</p>", db_path=temp_db)
    assert reserved is True
    assert status == "RESERVED"

    reserved2, status2 = reserve_draft_slot("F|A|1", "Cliente A", "a@demo.com", "Recordatorio", "<p>x</p>", db_path=temp_db)
    assert reserved2 is False
    assert status2 == "PROCESSING"

    mark_draft_failed("F|A|1", "Recordatorio", "falló", db_path=temp_db)
    reserved3, status3 = reserve_draft_slot("F|A|1", "Cliente A", "a@demo.com", "Recordatorio", "<p>x</p>", db_path=temp_db)
    assert reserved3 is True
    assert status3 == "RESERVED"

    mark_draft_created("F|A|1", "Recordatorio", "MSG-1", "<p>ok</p>", db_path=temp_db)
    assert draft_already_exists("F|A|1", "Recordatorio", db_path=temp_db)


def test_scheduler_ignores_excel_temp_files_and_continues_after_faulty(monkeypatch, tmp_path):
    # Arrange
    folder = tmp_path / "inbox"
    folder.mkdir()
    ok_file = folder / "ok.xlsx"
    bad_file = folder / "bad.xlsx"
    temp_excel = folder / "~$lock.xlsx"
    ok_file.write_bytes(b"ok")
    bad_file.write_bytes(b"bad")
    temp_excel.write_bytes(b"temp")

    import services
    calls = []

    def fake_import(path, db_path=None):
        calls.append(Path(path).name)
        if Path(path).name == "bad.xlsx":
            raise ValueError("bad file")
        return {"status": "OK"}

    monkeypatch.setattr(services, "import_signo_file", fake_import)
    cfg = {"signo_folder": str(folder), "days_advance": 7}

    # Act
    _check_folder(cfg)

    # Assert
    assert "~$lock.xlsx" not in calls
    assert "ok.xlsx" in calls
    assert "bad.xlsx" in calls
