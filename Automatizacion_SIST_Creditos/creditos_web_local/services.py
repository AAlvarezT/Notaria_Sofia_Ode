from __future__ import annotations

import hashlib
import json
import re
import sqlite3
import unicodedata
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook

from database import connect, get_db_path


def normalize_name(value: Any) -> str:
    text = str(value or "").strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.upper()
    text = re.sub(r"[^A-Z0-9]+", "", text)
    return text


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return re.sub(r"\s+", " ", text)


def _cell_text_preserving_zeros(value: Any) -> str:
    """Return string without stripping leading zeros.

    openpyxl returns text cells as str and numeric cells as int/float.
    For integer-valued floats (e.g. 30712.0) we drop the decimal part.
    True strings (e.g. "0030712") are returned unchanged.
    """
    if value is None:
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


def parse_decimal(value: Any) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).strip().replace("S/", "").replace("US$", "").replace("$", "")
    text = text.replace(" ", "")
    if text in ("", "-"):
        return Decimal("0")
    if text.count(",") and text.count("."):
        text = text.replace(".", "").replace(",", ".")
    elif text.count(","):
        text = text.replace(".", "").replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"Monto no válido: {value!r}") from exc


def parse_date(value: Any) -> Optional[str]:
    """Return ISO date string (YYYY-MM-DD), or None for absent/invalid dates."""
    if value in (None, "", "0000-00-00", "0000/00/00"):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, str):
        text = clean_text(value)
        if not text or text.startswith("0000"):
            return None
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d"):
            try:
                return datetime.strptime(text, fmt).date().isoformat()
            except ValueError:
                continue
        return text
    try:
        return value.isoformat()
    except AttributeError:
        return str(value)


def parse_timestamp(value: Any) -> Optional[str]:
    """Return ISO 8601 datetime string including time for version comparison."""
    if value in (None, "", "0000-00-00", "0000/00/00"):
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%dT%H:%M:%S")
    if isinstance(value, str):
        text = clean_text(value)
        if not text or text.startswith("0000"):
            return None
        for fmt in (
            "%Y-%m-%dT%H:%M:%S",
            "%Y-%m-%d %H:%M:%S",
            "%d/%m/%Y %H:%M:%S",
            "%d/%m/%Y %H:%M",
            "%d/%m/%Y",
            "%Y-%m-%d",
            "%d-%m-%Y",
            "%Y/%m/%d",
        ):
            try:
                return datetime.strptime(text, fmt).strftime("%Y-%m-%dT%H:%M:%S")
            except ValueError:
                continue
        return text
    return None


def document_key(tipo: Any, serie: Any, numero: Any) -> str:
    t = clean_text(tipo).upper()
    s = clean_text(serie).upper()
    n = clean_text(numero)
    return f"{t}|{s}|{n}"


def file_hash(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


HEADER_ALIASES: Dict[str, str] = {
    "DOI": "DOI",
    "RUC": "DOI",
    "RUCDNI": "DOI",
    "CLIENTE": "CLIENTE",
    "GRUPO": "CLIENTE",
    "FECHAEMISION": "FECHAEMISION",
    "EMISION": "FECHAEMISION",
    "TIPO": "TIPO",
    "SERIE": "SERIE",
    "SERIEA": "SERIE",
    "SERIES": "SERIE",
    "NUMERO": "NUMERO",
    "DOC": "NUMERO",
    "DOCUMENTO": "NUMERO",
    "KARDEX": "KARDEX",
    "TRAMITE": "KARDEX",
    "IMPORTEFACTURADO": "IMPORTEFACTURADO",
    "FACTURADO": "IMPORTEFACTURADO",
    "SALDOPENDIENTE": "SALDOPENDIENTE",
    "SALDO": "SALDOPENDIENTE",
    "FORMAPAGO": "FORMAPAGO",
    "PAGO": "FORMAPAGO",
    "FECHAVENCIMIENTO": "FECHAVENCIMIENTO",
    "VENCIMIENTO": "FECHAVENCIMIENTO",
    "EMPLEADO": "EMPLEADO",
    "FECHAACTUALIZACION": "FECHAACTUALIZACION",
    "ACTUALIZACION": "FECHAACTUALIZACION",
    "FECHACANCELACION": "FECHACANCELACION",
    "CANCELACION": "FECHACANCELACION",
}

# Alternate sheet names → canonical name
_SHEET_NAME_CANONICAL: Dict[str, str] = {
    "VENTAS": "VENTAS",
    "PENDIENTES": "PENDIENTES",
    "CANCELADOS": "CANCELADOS",
    "RECUPERADOHOJA1": "PENDIENTES",
    "HOJARECUPERADA": "PENDIENTES",
}


def _normalize_header(value: Any) -> str:
    return normalize_name(value).replace("\n", "")


def _canonical_header(value: Any) -> str:
    header = _normalize_header(value)
    return HEADER_ALIASES.get(header, header)


# ---------------------------------------------------------------------------
# Sheet reading with iter_rows (efficient for read_only workbooks)
# ---------------------------------------------------------------------------

def _read_sheet(ws, sheet_name: str) -> Tuple[List[Dict[str, Any]], List[str]]:
    """Parse one SIGNO sheet into a list of record dicts.

    Handles:
    - Title row in position 1 + header row in position 2 (flexible: scans up to row 40).
    - Leading zeros in DOI, serie and numero.
    - Proper ISO date for fecha_emision / fecha_vencimiento / fecha_cancelacion.
    - ISO timestamp for fecha_actualizacion (preserves hour for version comparison).
    - Invalid amounts flagged instead of silently set to zero.
    """
    canonical = _SHEET_NAME_CANONICAL.get(normalize_name(sheet_name))
    if canonical is None:
        return [], []

    # Load all rows at once – efficient even with read_only mode
    all_row_values: List[tuple] = list(ws.iter_rows(values_only=True))

    # Detect header row: first row (within first 40) that contains TIPO, SERIE, NUMERO
    required = {"TIPO", "SERIE", "NUMERO"}
    header_idx: Optional[int] = None
    col_map: Dict[str, int] = {}  # canonical header → 0-based column index

    for idx, row_vals in enumerate(all_row_values[:40]):
        found: Dict[str, int] = {}
        for col_i, cell_val in enumerate(row_vals):
            canon = _canonical_header(cell_val)
            if canon in required:
                found[canon] = col_i
        if required.issubset(found):
            header_idx = idx
            for col_i, cell_val in enumerate(row_vals):
                canon = _canonical_header(cell_val)
                if canon:
                    col_map[canon] = col_i
            break

    if header_idx is None:
        return [], [f"Hoja {sheet_name!r}: sin cabecera reconocida"]

    parsed_rows: List[Dict[str, Any]] = []
    errors: List[str] = []

    for row_offset, row_vals in enumerate(all_row_values[header_idx + 1:], start=header_idx + 2):

        def get(key: str) -> Any:
            i = col_map.get(key)
            if i is None or i >= len(row_vals):
                return None
            return row_vals[i]

        tipo_raw = get("TIPO")
        serie_raw = get("SERIE")
        numero_raw = get("NUMERO")

        # Skip completely empty rows
        if all(v in (None, "") for v in (tipo_raw, serie_raw, numero_raw)):
            continue

        tipo = clean_text(tipo_raw)
        serie = _cell_text_preserving_zeros(serie_raw)
        numero = _cell_text_preserving_zeros(numero_raw)

        dk = document_key(tipo, serie, numero)
        parts = dk.split("|")
        if len(parts) != 3 or all(p == "" for p in parts):
            errors.append(f"{sheet_name}: fila {row_offset} sin clave de comprobante válida")
            continue

        record: Dict[str, Any] = {
            "document_key": dk,
            "tipo": tipo,
            "serie": serie,
            "numero": numero,
            "source_sheet": sheet_name,
            "doi": _cell_text_preserving_zeros(get("DOI")),
            "cliente": clean_text(get("CLIENTE")),
            "kardex": clean_text(get("KARDEX")),
            "forma_pago": clean_text(get("FORMAPAGO")),
            "empleado": clean_text(get("EMPLEADO")),
            "fecha_emision": parse_date(get("FECHAEMISION")),
            "fecha_vencimiento": parse_date(get("FECHAVENCIMIENTO")),
            "fecha_cancelacion": parse_date(get("FECHACANCELACION")),
            "fecha_actualizacion": parse_timestamp(get("FECHAACTUALIZACION")),
        }

        has_invalid_amount = False
        for excel_col, db_field in (("IMPORTEFACTURADO", "importe_facturado"), ("SALDOPENDIENTE", "saldo_pendiente")):
            raw_val = get(excel_col)
            try:
                record[db_field] = str(parse_decimal(raw_val))
            except ValueError:
                errors.append(f"{sheet_name}: fila {row_offset} monto inválido en {excel_col} ({raw_val!r})")
                record[db_field] = None  # None = invalid; prevents silent "paid" classification
                has_invalid_amount = True

        if has_invalid_amount:
            record["_has_invalid_amount"] = True

        parsed_rows.append(record)

    return parsed_rows, errors


def _document_sort_key(record: Dict[str, Any]) -> tuple:
    sheet_priority = {"PENDIENTES": 3, "VENTAS": 2, "CANCELADOS": 1}
    sheet_name = (record.get("source_sheet") or "").upper()
    # For multi-sheet merged records source_sheet may be "CANCELADOS, PENDIENTES, VENTAS"
    priority = max((sheet_priority.get(s.strip(), 0) for s in sheet_name.split(",")), default=0)
    fecha = record.get("fecha_actualizacion") or "0000-00-00"
    try:
        saldo = Decimal(str(record.get("saldo_pendiente") or "0"))
    except InvalidOperation:
        saldo = Decimal("0")
    return (priority, fecha, saldo)


def _safe_decimal(value: Any) -> Optional[Decimal]:
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError):
        return None


def _merge_document_rows(rows: List[Dict[str, Any]]) -> Dict[str, Any]:
    """Merge rows from multiple sheets into one canonical record.

    Business rules applied:
    - When PENDIENTES is present, use its saldo as the authoritative balance.
    - Take the highest-priority (PENDIENTES > VENTAS > CANCELADOS) row as base.
    - For fields missing in the base row, fill from lower-priority rows.
    - Propagate _has_invalid_amount flag.
    """
    if not rows:
        return {}

    source_sheets: set = {(r.get("source_sheet") or "").upper() for r in rows}
    in_pendientes = "PENDIENTES" in source_sheets

    rows_sorted = sorted(rows, key=_document_sort_key, reverse=True)
    merged = dict(rows_sorted[0])

    for row in rows_sorted[1:]:
        for key in (
            "tipo", "serie", "numero", "doi", "cliente", "fecha_emision",
            "forma_pago", "fecha_vencimiento", "empleado", "fecha_actualizacion",
            "fecha_cancelacion", "kardex",
        ):
            value = row.get(key)
            if value in (None, ""):
                continue
            existing = merged.get(key)
            if existing in (None, ""):
                merged[key] = value
            elif key == "fecha_actualizacion" and str(value) > str(existing):
                merged[key] = value
            elif key == "cliente" and len(str(value)) > len(str(existing)):
                merged[key] = value

    # Saldo: PENDIENTES is the authoritative source when present
    if in_pendientes:
        pendientes_rows = [r for r in rows if "PENDIENTES" in (r.get("source_sheet") or "").upper()]
        best_pendiente = max(pendientes_rows, key=_document_sort_key)
        merged["saldo_pendiente"] = best_pendiente.get("saldo_pendiente")
        imp = best_pendiente.get("importe_facturado")
        if imp not in (None, ""):
            merged["importe_facturado"] = imp
        # If later rows from other sheets contradict Pendientes and no safe resolution
        # is possible, mark the document for manual review.
        pend_ts = best_pendiente.get("fecha_actualizacion") or "0000-00-00"
        pend_saldo = _safe_decimal(best_pendiente.get("saldo_pendiente"))
        contradictory_later = False
        for other in rows:
            other_sheet = (other.get("source_sheet") or "").upper()
            if "PENDIENTES" in other_sheet:
                continue
            other_ts = other.get("fecha_actualizacion") or "0000-00-00"
            if str(other_ts) <= str(pend_ts):
                continue
            other_saldo = _safe_decimal(other.get("saldo_pendiente"))
            if pend_saldo is not None and other_saldo is not None and other_saldo != pend_saldo:
                contradictory_later = True
                break

        # Multiple pendientes with same latest timestamp but different balances are
        # also contradictory.
        latest_pend_ts = max((r.get("fecha_actualizacion") or "0000-00-00" for r in pendientes_rows), default="0000-00-00")
        latest_pend_saldos = {
            _safe_decimal(r.get("saldo_pendiente"))
            for r in pendientes_rows
            if (r.get("fecha_actualizacion") or "0000-00-00") == latest_pend_ts
        }
        latest_pend_saldos.discard(None)
        if contradictory_later or len(latest_pend_saldos) > 1:
            merged["_has_conflict"] = True
    else:
        # No PENDIENTES: use max reported importe and saldo from available sheets
        best_imp: Optional[str] = None
        best_sal: Optional[str] = None
        for row in rows:
            for field, best_ref in (("importe_facturado", "best_imp"), ("saldo_pendiente", "best_sal")):
                val = row.get(field)
                if val is None:
                    continue
                try:
                    current_best = best_imp if field == "importe_facturado" else best_sal
                    if current_best is None or Decimal(str(val)) > Decimal(str(current_best)):
                        if field == "importe_facturado":
                            best_imp = val
                        else:
                            best_sal = val
                except InvalidOperation:
                    pass
        if best_imp is not None:
            merged["importe_facturado"] = best_imp
        if best_sal is not None:
            merged["saldo_pendiente"] = best_sal

    if any(r.get("_has_invalid_amount") for r in rows):
        merged["_has_invalid_amount"] = True

    merged["source_sheet"] = ", ".join(sorted(source_sheets))
    return merged


def _is_newer_report(existing: Optional[sqlite3.Row], candidate: Dict[str, Any]) -> bool:
    """True when candidate report is the same age or more recent than what is stored."""
    if existing is None:
        return True
    existing_ts = existing["fecha_actualizacion"] if existing["fecha_actualizacion"] else "0000-00-00"
    candidate_ts = candidate.get("fecha_actualizacion") or "0000-00-00"
    return str(candidate_ts) >= str(existing_ts)


def import_signo_file(path: str | Path, db_path: Optional[str | Path] = None) -> Dict[str, Any]:
    excel_path = Path(path)
    file_digest = file_hash(excel_path)
    conn = sqlite3.connect(str(db_path or get_db_path()))
    conn.row_factory = sqlite3.Row

    try:
        workbook = load_workbook(excel_path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"No se pudo abrir el archivo Excel: {exc}") from exc

    all_rows: List[Dict[str, Any]] = []
    errors: List[str] = []
    sheet_counts: Dict[str, int] = {}

    for sheet_name in workbook.sheetnames:
        canonical = _SHEET_NAME_CANONICAL.get(normalize_name(sheet_name))
        if canonical is None:
            continue
        ws = workbook[sheet_name]
        rows, sheet_errors = _read_sheet(ws, sheet_name)
        sheet_counts[canonical] = sheet_counts.get(canonical, 0) + len(rows)
        all_rows.extend(rows)
        errors.extend(sheet_errors)
    workbook.close()

    # Group rows by document identity
    doc_rows: Dict[str, List[Dict[str, Any]]] = {}
    for row in all_rows:
        doc_rows.setdefault(row["document_key"], []).append(row)

    import_at = datetime.now().strftime("%Y-%m-%dT%H:%M:%S")
    import_id = conn.execute(
        "INSERT INTO imports(file_name, file_hash, file_path, imported_at, status, errors, rows_seen, rows_saved)"
        " VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
        (
            excel_path.name, file_digest, str(excel_path), import_at,
            "OK" if not errors else "WARN",
            json.dumps(errors, ensure_ascii=False),
            len(all_rows), len(doc_rows),
        ),
    ).lastrowid

    saved = 0
    for doc_key in sorted(doc_rows):
        row = _merge_document_rows(doc_rows[doc_key])
        existing = conn.execute("SELECT * FROM documents WHERE document_key = ?", (doc_key,)).fetchone()

        if existing is not None and not _is_newer_report(existing, row):
            continue

        state = determine_document_state(doc_key, row)
        saldo_val = row.get("saldo_pendiente")
        saldo_str = str(saldo_val) if saldo_val is not None else "0"
        importe_val = row.get("importe_facturado")
        importe_str = str(importe_val) if importe_val is not None else "0"

        doc = {
            "document_key": doc_key,
            "tipo": row.get("tipo", ""),
            "serie": row.get("serie", ""),
            "numero": row.get("numero", ""),
            "doi": row.get("doi", ""),
            "cliente": row.get("cliente", ""),
            "fecha_emision": row.get("fecha_emision"),
            "importe_facturado": importe_str,
            "saldo_observado": saldo_str,
            "forma_pago": row.get("forma_pago", ""),
            "fecha_vencimiento": row.get("fecha_vencimiento"),
            "empleado": row.get("empleado", ""),
            "fecha_actualizacion": row.get("fecha_actualizacion"),
            "estado": state,
            "fuente": row.get("source_sheet", ""),
            "ultima_importacion_id": import_id,
            # SALDO_PARCIAL is a valid business state, not an inconsistency
            "revision_required": 1 if state == "EN_REVISION" else 0,
            # Preserve manually entered data
            "observaciones": existing["observaciones"] if existing is not None else "",
            "created_at": existing["created_at"] if existing is not None else import_at,
            "updated_at": import_at,
        }
        conn.execute(
            """INSERT INTO documents(
                document_key, tipo, serie, numero, doi, cliente, fecha_emision,
                importe_facturado, saldo_observado, forma_pago, fecha_vencimiento,
                empleado, fecha_actualizacion, estado, fuente, ultima_importacion_id,
                revision_required, observaciones, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(document_key) DO UPDATE SET
                tipo = excluded.tipo,
                serie = excluded.serie,
                numero = excluded.numero,
                doi = excluded.doi,
                cliente = excluded.cliente,
                fecha_emision = excluded.fecha_emision,
                importe_facturado = excluded.importe_facturado,
                saldo_observado = excluded.saldo_observado,
                forma_pago = excluded.forma_pago,
                fecha_vencimiento = CASE
                    WHEN documents.fecha_vencimiento IS NOT NULL AND documents.fecha_vencimiento <> ''
                    THEN documents.fecha_vencimiento
                    ELSE excluded.fecha_vencimiento
                END,
                empleado = excluded.empleado,
                fecha_actualizacion = excluded.fecha_actualizacion,
                estado = excluded.estado,
                fuente = excluded.fuente,
                ultima_importacion_id = excluded.ultima_importacion_id,
                revision_required = excluded.revision_required,
                observaciones = CASE
                    WHEN documents.observaciones IS NOT NULL AND documents.observaciones <> ''
                    THEN documents.observaciones
                    ELSE excluded.observaciones
                END,
                updated_at = excluded.updated_at""",
            (
                doc["document_key"], doc["tipo"], doc["serie"], doc["numero"],
                doc["doi"], doc["cliente"], doc["fecha_emision"],
                doc["importe_facturado"], doc["saldo_observado"], doc["forma_pago"],
                doc["fecha_vencimiento"], doc["empleado"], doc["fecha_actualizacion"],
                doc["estado"], doc["fuente"], doc["ultima_importacion_id"],
                doc["revision_required"], doc["observaciones"],
                doc["created_at"], doc["updated_at"],
            ),
        )
        saved += 1

    conn.commit()
    conn.close()

    return {
        "import_id": import_id,
        "rows_seen": len(all_rows),
        "documents": saved,
        "sheet_counts": sheet_counts,
        "errors": errors,
        "status": "OK" if not errors else "WARN",
    }


def determine_document_state(doc_key: str, row: Dict[str, Any]) -> str:
    """Classify a document into its collection state.

    States:
    - EXCLUIDO     : saldo ≤ 0 (paid or zero)
    - SALDO_PARCIAL: partially paid; saldo > 0 but < importe (valid, shows in cobranza)
    - PENDIENTE    : full balance owed, confirmed by PENDIENTES sheet
    - EN_REVISION  : data inconsistency or source unknown – requires human review
    """
    if row.get("_has_invalid_amount"):
        return "EN_REVISION"
    if row.get("_has_conflict"):
        return "EN_REVISION"

    saldo_raw = row.get("saldo_pendiente")
    if saldo_raw is None:
        return "EN_REVISION"

    try:
        saldo = Decimal(str(saldo_raw))
    except InvalidOperation:
        return "EN_REVISION"

    if saldo <= 0:
        return "EXCLUIDO"

    importe_raw = row.get("importe_facturado")
    try:
        importe = Decimal(str(importe_raw)) if importe_raw is not None else Decimal("0")
    except InvalidOperation:
        importe = Decimal("0")

    if importe > 0 and saldo < importe:
        return "SALDO_PARCIAL"

    source = (row.get("source_sheet") or "").upper()
    if "PENDIENTES" in source:
        return "PENDIENTE"

    return "EN_REVISION"


def list_reviews() -> List[Dict[str, Any]]:
    with connect() as conn:
        rows = conn.execute(
            "SELECT * FROM documents WHERE revision_required = 1 OR estado IN ('EN_REVISION', 'EXCLUIDO')"
            " ORDER BY updated_at DESC"
        ).fetchall()
    return [dict(row) for row in rows]


def list_clients() -> List[Dict[str, Any]]:
    with connect() as conn:
        rows = conn.execute("SELECT DISTINCT cliente, doi FROM documents ORDER BY cliente").fetchall()
    return [dict(row) for row in rows]


def create_review_item(document_key: str, kind: str, detail: str) -> None:
    with connect() as conn:
        conn.execute(
            "INSERT OR IGNORE INTO review_queue(document_key, kind, detail, created_at)"
            " VALUES (?, ?, ?, datetime('now'))",
            (document_key, kind, detail),
        )
